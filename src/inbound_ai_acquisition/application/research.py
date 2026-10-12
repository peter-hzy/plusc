"""真实获客研究的可恢复纵向工作流。"""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from ..adapters.registry import built_in_adapters
from ..domain.company_fit import SERVICE_FIELDS, segment_allowed, validate_service_fit
from ..domain.policies import require_one_cta
from ..errors import StateError, ValidationError
from ..storage.database import Database, canonical_hash, new_id, utc_now
from ..storage.repositories import active_stop_exists
from .discovery import discovery_coverage
from .services import create_draft, decide_claim, normalize_name, normalize_url, root_domain


def _field_evidence(connection: Any, record: dict[str, Any], field: str, evidence_ids: list[str]) -> list[str]:
    """Explicit URL bindings narrow evidence; legacy bundles retain review-only broad bindings."""
    bindings = record.get("field_sources")
    if bindings is None:
        return evidence_ids
    if not isinstance(bindings, dict):
        raise ValidationError("field_sources 必须是字段到URL数组的映射")
    urls = bindings.get(field)
    if not isinstance(urls, list) or not urls or any(not isinstance(u, str) for u in urls):
        raise ValidationError(f"field_sources 缺少{field}的来源")
    wanted = {normalize_url(u) for u in urls}
    rows = [
        connection.execute("SELECT source_url FROM evidence WHERE evidence_id=?", (eid,)).fetchone()
        for eid in evidence_ids
    ]
    selected = [eid for eid, row in zip(evidence_ids, rows, strict=True) if row[0] in wanted]
    if not selected or wanted - {row[0] for row in rows}:
        raise ValidationError(f"{field}关联了未在本实体成功读取的来源")
    return selected


_CANDIDATE_STATUSES = {"candidate_pool", "selected", "held", "excluded"}

_FIT_GATE_VALUES = {
    "source_market_role": {"outbound_seller", "channel_partner", "unclear", "not_source_market"},
    "china_product_signal": {"direct", "adjacent", "none"},
    "ground_partner_route": {"external_partner_possible", "unclear", "in_house"},
    "competitive_position": {"potential_buyer", "mixed", "competitor", "unknown"},
}


def validate_company_fit_gates(company: dict[str, Any]) -> dict[str, str]:
    """New records use service fit; legacy records retain their historical contract."""
    policy = company.get("qualification_policy")
    if policy == "service_fit_v1":
        return validate_service_fit(company)
    if policy is not None:
        raise ValidationError("未知 qualification_policy")
    name = str(company.get("name", "未命名公司")).strip()
    values: dict[str, str] = {}
    for field, allowed in _FIT_GATE_VALUES.items():
        value = str(company.get(field, "")).strip()
        if not value:
            labels = {
                "source_market_role": "客源市场角色",
                "china_product_signal": "中国产品信号",
                "ground_partner_route": "中国地接合作路径",
                "competitive_position": "竞争关系",
            }
            raise ValidationError(f"{name}：缺少{labels[field]}")
        if value not in allowed:
            raise ValidationError(f"{name}：{field} 值无效：{value}")
        values[field] = value
    path = str(company.get("qualification_path", "direct"))
    if path not in {"direct", "expansion", "channel"}:
        raise ValidationError("qualification_path 无效")
    if values["source_market_role"] != ("channel_partner" if path == "channel" else "outbound_seller"):
        raise ValidationError(f"{name}：客源市场角色不是向本市场客人销售出境产品的公司")
    if path == "direct" and values["china_product_signal"] == "none":
        raise ValidationError(f"{name}：没有中国产品或可解释的亚洲长线适配信号")
    if path != "direct":
        for field in ("audience_fit", "cooperation_hypothesis", "counter_evidence", "next_action"):
            if not str(company.get(field, "")).strip():
                raise ValidationError(f"{name}：{path}路径缺少{field}")
        if str(company.get("fit_status")) == "priority_chat":
            raise ValidationError(f"{name}：扩展/渠道假设仅进入待确认对话池，不自动升优先")
    status = str(company.get("fit_status", "")).strip()
    if status in {"priority_chat", "chat_candidate"} and values["competitive_position"] == "competitor":
        raise ValidationError(f"{name}：一体化竞争者不能进入优先或可联系名单")
    if status == "priority_chat" and values["ground_partner_route"] != "external_partner_possible":
        raise ValidationError(f"{name}：优先名单必须有外部中国地接合作路径证据")
    return values


def _candidate_group(candidate: dict[str, Any], default_query: str) -> tuple[str, str]:
    category = str(candidate.get("source_category") or "public_web_search").strip()
    query = str(candidate.get("query_summary") or default_query).strip()
    if not category or not query:
        raise ValidationError("候选公司必须记录 source_category 和 query_summary")
    return category, query


def _candidate_observations(candidate: dict[str, Any], default_query: str) -> list[dict[str, Any]]:
    primary = {
        "source_url": candidate.get("source_url"),
        "observed_text": candidate.get("observed_text"),
        "source_category": candidate.get("source_category", "public_web_search"),
        "query_summary": candidate.get("query_summary", default_query),
        "observed_at": candidate.get("observed_at"),
    }
    extras = candidate.get("sources", [])
    if extras and not isinstance(extras, list):
        raise ValidationError("候选公司的 sources 必须是数组")
    return [primary, *extras]


def _read_bundle(path: Path | None) -> dict[str, Any] | None:
    if path is None:
        return None
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValidationError("research bundle 必须是 JSON object")
    return payload


def _run_exists(connection: Any, run_id: str) -> None:
    if not connection.execute("SELECT 1 FROM runs WHERE run_id=?", (run_id,)).fetchone():
        raise ValidationError("找不到 run")


def _insert_source_run(
    connection: Any,
    *,
    run_id: str,
    adapter_id: str,
    query_summary: str,
    outcome: str,
    result_count: int,
    error_code: str = "",
    retryable: bool = False,
) -> str:
    source_run_id = new_id("src")
    now = utc_now()
    connection.execute(
        "INSERT INTO source_runs VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (
            source_run_id,
            run_id,
            adapter_id,
            query_summary,
            canonical_hash({"adapter": adapter_id, "query": query_summary}),
            now,
            now,
            1,
            outcome,
            result_count,
            0.0,
            0,
            error_code,
            int(retryable),
        ),
    )
    return source_run_id


def _insert_evidence(connection: Any, source_run_id: str, entity_type: str, entity_id: str, row: dict[str, Any]) -> str:
    evidence_id = str(row.get("evidence_id") or new_id("ev"))
    text = str(row["observed_text"]).strip()
    existing = connection.execute(
        "SELECT evidence_id FROM evidence WHERE entity_type=? AND entity_id=? AND source_url=? "
        "AND content_hash=? AND purpose=? AND access_status=? AND conflicts=? AND source_level=? "
        "AND published_at IS ? AND (?='' OR observed_at=?)",
        (
            entity_type,
            entity_id,
            normalize_url(str(row["source_url"])),
            canonical_hash(text),
            str(row.get("purpose", "")),
            str(row.get("access_status", "accessible")),
            int(bool(row.get("conflicts", False))),
            str(row.get("source_level", "secondary")),
            row.get("published_at"),
            str(row.get("observed_at") or ""),
            str(row.get("observed_at") or ""),
        ),
    ).fetchone()
    if existing:
        return existing["evidence_id"]
    connection.execute(
        "INSERT INTO evidence VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (
            evidence_id,
            source_run_id,
            entity_type,
            entity_id,
            normalize_url(str(row["source_url"])),
            str(row.get("source_level", "secondary")),
            str(row.get("purpose", "")),
            text,
            row.get("published_at"),
            str(row.get("observed_at") or utc_now()),
            canonical_hash(text),
            str(row.get("access_status", "accessible")),
            int(bool(row.get("conflicts", False))),
        ),
    )
    return evidence_id


def _insert_claim(
    connection: Any,
    *,
    entity_type: str,
    entity_id: str,
    field: str,
    value: str,
    evidence_ids: list[str],
    grade: str,
) -> str:
    # Replaying an observation must not duplicate or silently reapprove a fact.
    for previous in connection.execute(
        "SELECT claim_id,proposed_value FROM claims WHERE entity_type=? AND entity_id=? AND field_name=? "
        "AND status<>'superseded' ORDER BY rowid DESC LIMIT 1",
        (entity_type, entity_id, field),
    ).fetchall():
        linked = {
            r[0]
            for r in connection.execute(
                "SELECT evidence_id FROM claim_evidence WHERE claim_id=?", (previous["claim_id"],)
            )
        }
        if previous["proposed_value"] == value and linked == set(evidence_ids):
            return previous["claim_id"]
    claim_id = new_id("cl")
    connection.execute(
        "UPDATE claims SET status='superseded' WHERE entity_type=? AND entity_id=? AND field_name=? AND status='proposed'",
        (entity_type, entity_id, field),
    )
    connection.execute(
        "INSERT INTO claims VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (
            claim_id,
            entity_type,
            entity_id,
            field,
            value,
            "proposed",
            grade,
            "",
            None,
            None,
            "",
            None,
            utc_now(),
        ),
    )
    for evidence_id in evidence_ids:
        connection.execute("INSERT INTO claim_evidence VALUES (?,?)", (claim_id, evidence_id))
    return claim_id


def ingest_source_bundle(db: Database, run_id: str, bundle: dict[str, Any], request_id: str) -> dict[str, Any]:
    """一次性保存 Codex 已执行的发现、公司核验和联系人核验结果。"""
    if bundle.get("executor") != "codex_bridge":
        raise ValidationError("V1 真实研究 bundle 的 executor 必须是 codex_bridge")
    companies = bundle.get("companies")
    discovered_candidates = bundle.get("discovered_candidates")
    discovered_count = int(bundle.get("discovered_count", 0))
    if not isinstance(companies, list):
        raise ValidationError("research bundle 的 companies 必须是数组，可以为空")
    if not isinstance(discovered_candidates, list) or len(discovered_candidates) != discovered_count:
        raise ValidationError("discovered_count 必须与 discovered_candidates 的真实条目数一致")
    if discovered_count < len(companies):
        raise ValidationError("本轮入选公司不能多于本轮发现候选")
    bridge = built_in_adapters()["codex_bridge"]
    default_query = str(bundle.get("discovery_query", "Codex public-web discovery"))
    source_plan = bundle.get("source_plan", [])
    if source_plan and not isinstance(source_plan, list):
        raise ValidationError("source_plan 必须是数组")
    if not source_plan:
        source_plan = [
            {
                "source_category": "public_web_search",
                "source_role": "discover",
                "target_signal": "目标公司与公开业务信号",
                "query_or_entry_url": default_query,
                "coverage": str(bundle.get("coverage", "")),
                "access_status": "opened",
            }
        ]
    grouped_candidates: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for candidate in discovered_candidates:
        for observation in _candidate_observations(candidate, default_query):
            grouped_candidates.setdefault(_candidate_group(observation, default_query), []).append(observation)
    discovery_results = {
        key: bridge.discover({"query_summary": key[1], "outcome": "succeeded", "observations": rows})
        for key, rows in grouped_candidates.items()
    }
    candidate_by_name: dict[str, dict[str, Any]] = {}
    for candidate in discovered_candidates:
        name = normalize_name(str(candidate.get("name", "")))
        if not name or name in candidate_by_name:
            raise ValidationError("discovered_candidates 的公司名称不能为空或重复")
        candidate_by_name[name] = candidate
    changed = 0
    with db.transaction() as connection:
        _run_exists(connection, run_id)
        now = utc_now()
        for entry in source_plan:
            for required in ("source_category", "query_or_entry_url"):
                if not str(entry.get(required, "")).strip():
                    raise ValidationError(f"source_plan 缺少字段：{required}")
            cursor = connection.execute(
                "INSERT INTO source_plan_entries VALUES (?,?,?,?,?,?,?,?,?,?,?) "
                "ON CONFLICT(run_id,source_category,query_or_entry_url) DO UPDATE SET "
                "source_role=excluded.source_role,target_signal=excluded.target_signal,coverage=excluded.coverage,"
                "access_status=excluded.access_status,contribution=excluded.contribution,updated_at=excluded.updated_at",
                (
                    new_id("plan"),
                    run_id,
                    str(entry["source_category"]).strip(),
                    str(entry.get("source_role", "discover")).strip(),
                    str(entry.get("target_signal", "")).strip(),
                    str(entry["query_or_entry_url"]).strip(),
                    str(entry.get("coverage", "")).strip(),
                    str(entry.get("access_status", "planned")).strip(),
                    str(entry.get("contribution", "")).strip(),
                    now,
                    now,
                ),
            )
            changed += max(cursor.rowcount, 0)
        discovery_source_runs: dict[tuple[str, str], str] = {}
        for key, result in discovery_results.items():
            discovery_source_runs[key] = _insert_source_run(
                connection,
                run_id=run_id,
                adapter_id=result.adapter_id,
                query_summary=key[1],
                outcome=result.outcome,
                result_count=len(result.observations),
            )
            changed += 1
        candidate_ids: dict[str, str] = {}
        candidate_ids_by_domain: dict[str, str] = {}
        for candidate in discovered_candidates:
            normalized_name = normalize_name(str(candidate["name"]))
            candidate_url = normalize_url(str(candidate["source_url"]))
            # 目录/展会页面属于来源平台，不是被列出的公司官网。
            company_url = str(candidate.get("website", "")).strip()
            if not company_url and candidate.get("source_category", "public_web_search") in {
                "official_company_site",
                "public_web_search",
            }:
                company_url = candidate_url
            candidate_domain = root_domain(company_url) if company_url else ""
            status = str(candidate.get("discovery_status", "candidate_pool"))
            if status not in _CANDIDATE_STATUSES - {"selected"}:
                raise ValidationError(f"候选状态不合法：{status}")
            existing = connection.execute(
                "SELECT candidate_id,name_normalized,root_domain FROM candidates WHERE run_id=? "
                "AND (name_normalized=? OR (?<>'' AND root_domain=?)) ORDER BY first_seen_at LIMIT 1",
                (run_id, normalized_name, candidate_domain, candidate_domain),
            ).fetchone()
            if existing:
                candidate_id = existing["candidate_id"]
                cursor = connection.execute(
                    "UPDATE candidates SET country_hint=CASE WHEN ?<>'' THEN ? ELSE country_hint END,"
                    "segment_hint=CASE WHEN ?<>'' THEN ? ELSE segment_hint END,"
                    "decision_reason=CASE WHEN ?<>'' THEN ? ELSE decision_reason END,"
                    "root_domain=CASE WHEN root_domain='' THEN ? ELSE root_domain END,last_seen_at=? WHERE candidate_id=?",
                    (
                        str(candidate.get("country_hint", "")).strip(),
                        str(candidate.get("country_hint", "")).strip(),
                        str(candidate.get("segment_hint", "")).strip(),
                        str(candidate.get("segment_hint", "")).strip(),
                        str(candidate.get("decision_reason", "")).strip(),
                        str(candidate.get("decision_reason", "")).strip(),
                        candidate_domain,
                        now,
                        candidate_id,
                    ),
                )
                changed += max(cursor.rowcount, 0)
            else:
                candidate_id = new_id("cand")
                connection.execute(
                    "INSERT INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        candidate_id,
                        run_id,
                        str(candidate["name"]).strip(),
                        normalized_name,
                        candidate_domain,
                        candidate_url,
                        str(candidate.get("country_hint", "")).strip(),
                        str(candidate.get("segment_hint", "")).strip(),
                        status,
                        str(candidate.get("decision_reason", "")).strip(),
                        None,
                        now,
                        now,
                    ),
                )
                changed += 1
            candidate_ids[normalized_name] = candidate_id
            if candidate_domain:
                candidate_ids_by_domain[candidate_domain] = candidate_id
            for observation in _candidate_observations(candidate, default_query):
                category, query = _candidate_group(observation, default_query)
                source_run_id = discovery_source_runs[(category, query)]
                cursor = connection.execute(
                    "INSERT OR IGNORE INTO candidate_sources VALUES (?,?,?,?,?,?,?,?)",
                    (
                        new_id("csrc"),
                        run_id,
                        candidate_id,
                        source_run_id,
                        category,
                        normalize_url(str(observation["source_url"])),
                        str(observation["observed_text"]).strip(),
                        str(observation.get("observed_at") or now),
                    ),
                )
                changed += max(cursor.rowcount, 0)
        seen_domains: set[str] = set()
        created_companies = 0
        for company in companies:
            for field in ("name", "country", "segment", "website", "fit_reason", "fit_status", "sources", "contact"):
                if field not in company:
                    raise ValidationError(f"公司研究缺少字段：{field}")
            fit_gates = validate_company_fit_gates(company)
            if company.get("qualification_policy") == "service_fit_v1" and not isinstance(
                company.get("field_sources"), dict
            ):
                raise ValidationError("新研究必须逐字段绑定来源 field_sources")
            website = normalize_url(str(company["website"]))
            domain = root_domain(website)
            if not domain or domain in seen_domains:
                raise ValidationError(f"公司官网域名缺失或重复：{domain}")
            seen_domains.add(domain)
            existing_company = connection.execute(
                "SELECT company_id FROM companies WHERE run_id=? AND (root_domain=? OR company_id IN "
                "(SELECT selected_company_id FROM candidates WHERE run_id=? AND candidate_id=?))",
                (
                    run_id,
                    domain,
                    run_id,
                    candidate_ids.get(normalize_name(str(company["name"]))) or candidate_ids_by_domain.get(domain),
                ),
            ).fetchone()
            if existing_company:
                candidate_id = candidate_ids.get(normalize_name(str(company["name"]))) or candidate_ids_by_domain.get(
                    domain
                )
                if candidate_id:
                    connection.execute(
                        "UPDATE candidates SET discovery_status='selected',selected_company_id=?,last_seen_at=? "
                        "WHERE candidate_id=?",
                        (existing_company["company_id"], now, candidate_id),
                    )
            company_id = existing_company["company_id"] if existing_company else new_id("co")
            if not existing_company:
                connection.execute(
                    "INSERT INTO companies VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        company_id,
                        run_id,
                        str(company["name"]).strip(),
                        normalize_name(str(company["name"])),
                        "",
                        str(company.get("city", "")),
                        "",
                        "",
                        "",
                        "",
                        "research_needed",
                        "",
                        str(company.get("risk_notes", "")),
                        now,
                        now,
                    ),
                )
                changed += 1
                created_companies += 1
            candidate = candidate_by_name.get(normalize_name(str(company["name"])))
            if candidate is None:
                candidate = next(
                    (
                        row
                        for row in discovered_candidates
                        if root_domain(normalize_url(str(row["source_url"]))) == domain
                    ),
                    None,
                )
            if candidate is None:
                raise ValidationError(f"最终公司不在主动发现清单中：{company['name']}")
            normalized_company = normalize_name(str(company["name"]))
            selected_candidate_id = candidate_ids.get(normalized_company) or candidate_ids_by_domain.get(domain)
            connection.execute(
                "UPDATE candidates SET discovery_status='selected',decision_reason=?,selected_company_id=?,last_seen_at=? "
                "WHERE candidate_id=?",
                (
                    str(candidate.get("decision_reason") or company.get("fit_reason", "")).strip(),
                    company_id,
                    now,
                    selected_candidate_id,
                ),
            )
            discovery_source_run_id = discovery_source_runs[_candidate_group(candidate, default_query)]
            _insert_evidence(
                connection,
                discovery_source_run_id,
                "company",
                company_id,
                {
                    **candidate,
                    "purpose": "discovery",
                    "source_level": str(candidate.get("source_level", "discovery")),
                },
            )
            changed += 1
            company_evidence: list[str] = []
            sources = company["sources"]
            if not isinstance(sources, list) or not sources:
                raise ValidationError("每家公司至少需要一条打开页面后的来源")
            for source in sources:
                operation = str(source.get("operation", "inspect_company"))
                if operation not in bridge.capabilities:
                    raise ValidationError(f"codex_bridge 不支持操作：{operation}")
                observation = {
                    "source_url": source.get("source_url"),
                    "observed_text": source.get("observed_text"),
                }
                result = getattr(bridge, operation)(
                    {
                        "query_summary": source.get("query_summary", ""),
                        "outcome": source.get("outcome", "succeeded"),
                        "observations": [observation],
                    }
                )
                source_run_id = _insert_source_run(
                    connection,
                    run_id=run_id,
                    adapter_id=result.adapter_id,
                    query_summary=str(source.get("query_summary", operation)),
                    outcome=result.outcome,
                    result_count=len(result.observations),
                    error_code=result.error_code,
                    retryable=result.retryable,
                )
                if result.outcome == "succeeded":
                    company_evidence.append(_insert_evidence(connection, source_run_id, "company", company_id, source))
                    changed += 2
            if not company_evidence:
                raise ValidationError(f"公司没有可用 Evidence：{company['name']}")
            for field, value, grade in (
                ("website", website, "primary"),
                ("country", str(company["country"]), "strong"),
                ("segment", str(company["segment"]), "strong"),
                ("fit_reason", str(company["fit_reason"]), "strong"),
                ("fit_status", str(company["fit_status"]), "strong"),
                *((field, value, "strong") for field, value in fit_gates.items()),
            ):
                _insert_claim(
                    connection,
                    entity_type="company",
                    entity_id=company_id,
                    field=field,
                    value=value,
                    evidence_ids=_field_evidence(connection, company, field, company_evidence),
                    grade=grade,
                )
                changed += 1
            for field in (
                "qualification_policy",
                "qualification_path",
                "audience_fit",
                "cooperation_hypothesis",
                "counter_evidence",
                "next_action",
                "group_key",
                "independent_unit_reason",
            ):
                if field not in fit_gates and str(company.get(field, "")).strip():
                    _insert_claim(
                        connection,
                        entity_type="company",
                        entity_id=company_id,
                        field=field,
                        value=str(company[field]),
                        evidence_ids=_field_evidence(connection, company, field, company_evidence),
                        grade="strong",
                    )
            contact = company["contact"]
            if not isinstance(contact, dict):
                raise ValidationError("contact 必须是 object")
            contact_sources = contact.get("sources", [])
            if not contact_sources:
                raise ValidationError("每家公司必须保存负责人或公开转交入口的来源")
            # Reuse the same observed route, not an arbitrary person at the company.
            route_key = tuple(
                str(contact.get(k, "")).strip().casefold() for k in ("public_email", "linkedin_url", "full_name")
            )
            existing_contact = None
            for stored in connection.execute("SELECT * FROM contacts WHERE company_id=?", (company_id,)).fetchall():
                preview = dict(stored)
                for fact in connection.execute(
                    "SELECT field_name,proposed_value FROM claims WHERE entity_type='contact' AND entity_id=? "
                    "AND status IN ('proposed','accepted') ORDER BY CASE status WHEN 'accepted' THEN 0 ELSE 1 END,rowid",
                    (stored["contact_id"],),
                ):
                    preview[fact["field_name"]] = fact["proposed_value"]
                if (
                    tuple(
                        str(preview.get(k, "")).strip().casefold()
                        for k in ("public_email", "linkedin_url", "full_name")
                    )
                    == route_key
                ):
                    existing_contact = stored
                    break
            contact_id = existing_contact["contact_id"] if existing_contact else new_id("ct")
            if not existing_contact:
                connection.execute(
                    "INSERT INTO contacts VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        contact_id,
                        company_id,
                        "",
                        "",
                        "",
                        str(contact.get("contact_role", "primary")),
                        "",
                        "",
                        "unverified",
                        "not_contacted",
                        None,
                        now,
                        now,
                    ),
                )
                changed += 1
            contact_evidence: list[str] = []
            for source in contact_sources:
                outcome = str(source.get("outcome", "succeeded"))
                observations = []
                if outcome == "succeeded":
                    observations = [
                        {"source_url": source.get("source_url"), "observed_text": source.get("observed_text")}
                    ]
                result = bridge.inspect_contact(
                    {
                        "query_summary": source.get("query_summary", "contact research"),
                        "outcome": outcome,
                        "error_code": source.get("error_code", ""),
                        "observations": observations,
                    }
                )
                source_run_id = _insert_source_run(
                    connection,
                    run_id=run_id,
                    adapter_id=result.adapter_id,
                    query_summary=str(source.get("query_summary", "contact research")),
                    outcome=result.outcome,
                    result_count=len(result.observations),
                    error_code=result.error_code,
                    retryable=result.retryable,
                )
                changed += 1
                if result.outcome == "succeeded":
                    contact_evidence.append(_insert_evidence(connection, source_run_id, "contact", contact_id, source))
                    changed += 1
            identity_status = str(contact.get("identity_status", "not_found"))
            if identity_status in {"supported", "verified"} and not contact_evidence:
                raise ValidationError("supported/verified 联系人必须有公开身份 Evidence")
            claim_fields = {
                "full_name": str(contact.get("full_name", "")),
                "title": str(contact.get("title", "")),
                "role_scope": str(contact.get("why_contact", contact.get("role_scope", ""))),
                "linkedin_url": str(contact.get("linkedin_url", "")),
                "public_email": str(contact.get("public_email", "")),
                "identity_status": identity_status,
            }
            route_grade = str(contact.get("route_grade", "R3"))
            route_purpose = str(contact.get("route_purpose", "unknown"))
            if route_grade not in {"R0", "R1", "R2", "R3"} or route_purpose not in {
                "business",
                "general",
                "unknown",
                "unsuitable",
            }:
                raise ValidationError("入口等级/用途无效")
            if route_grade in {"R1", "R2"} and (
                route_purpose != "business"
                or not claim_fields["role_scope"]
                or (route_grade == "R1" and not claim_fields["full_name"])
            ):
                raise ValidationError("R1/R2需要明确商务用途、职责；R1还需要公开姓名")
            claim_fields.update(route_grade=route_grade, route_purpose=route_purpose)
            for field, value in claim_fields.items():
                if value and contact_evidence:
                    _insert_claim(
                        connection,
                        entity_type="contact",
                        entity_id=contact_id,
                        field=field,
                        value=value,
                        evidence_ids=_field_evidence(connection, contact, field, contact_evidence),
                        grade="primary" if identity_status == "verified" else "strong",
                    )
                    changed += 1
            if identity_status not in {"supported", "verified"}:
                if not connection.execute(
                    "SELECT 1 FROM research_gaps WHERE entity_type='contact' AND entity_id=? "
                    "AND field_name='responsible_contact' AND status='open'",
                    (contact_id,),
                ).fetchone():
                    connection.execute(
                        "INSERT INTO research_gaps VALUES (?,?,?,?,?,?,?,?,?)",
                        (
                            new_id("gap"),
                            "contact",
                            contact_id,
                            "responsible_contact",
                            str(contact.get("gap_note", "已搜索但未找到可核验负责人或公开转交入口")),
                            1,
                            "open",
                            now,
                            None,
                        ),
                    )
                    changed += 1
        if connection.execute(
            "SELECT 1 FROM claims WHERE status='proposed' AND ((entity_type='company' AND entity_id IN "
            "(SELECT company_id FROM companies WHERE run_id=?)) OR (entity_type='contact' AND entity_id IN "
            "(SELECT contact_id FROM contacts JOIN companies USING(company_id) WHERE run_id=?)))",
            (run_id, run_id),
        ).fetchone():
            connection.execute(
                "UPDATE runs SET status='reviewing',next_action='awaiting_claim_review' WHERE run_id=?", (run_id,)
            )
        db.append_event(
            connection,
            event_type="research.sources_ingested",
            request_id=request_id,
            run_id=run_id,
            payload={
                "discovered_count": discovered_count,
                "selected_count": len(companies),
                "source_categories": sorted({key[0] for key in grouped_candidates}),
                "discovered_candidates": [
                    {
                        "name": row["name"],
                        "source_url": row["source_url"],
                        "source_category": _candidate_group(row, default_query)[0],
                        "discovery_status": "selected"
                        if normalize_name(str(row["name"])) in {normalize_name(str(item["name"])) for item in companies}
                        else str(row.get("discovery_status", "candidate_pool")),
                    }
                    for row in discovered_candidates
                ],
            },
        )
    coverage = discovery_coverage(db, run_id)
    next_action = db.query("SELECT next_action FROM runs WHERE run_id=?", (run_id,))[0]["next_action"]
    return {
        "status": "needs_review",
        "run_id": run_id,
        "changed_records": changed,
        "next_action": next_action,
        "coverage": coverage,
    }


def approve_run_claims(db: Database, run_id: str, actor: str, reason: str, request_id: str) -> dict[str, Any]:
    with db.transaction():
        return _approve_run_claims(db, run_id, actor, reason, request_id)


def _approve_run_claims(db: Database, run_id: str, actor: str, reason: str, request_id: str) -> dict[str, Any]:
    if not actor.strip() or not reason.strip():
        raise ValidationError("批次审核必须提供真实审核人和理由")
    claims = db.query(
        "SELECT cl.claim_id FROM claims cl LEFT JOIN companies co ON cl.entity_type='company' AND cl.entity_id=co.company_id "
        "LEFT JOIN contacts ct ON cl.entity_type='contact' AND cl.entity_id=ct.contact_id "
        "LEFT JOIN companies cc ON ct.company_id=cc.company_id "
        "WHERE cl.status='proposed' AND COALESCE(co.run_id,cc.run_id)=? ORDER BY cl.created_at,cl.rowid",
        (run_id,),
    )
    if not claims:
        raise StateError("该 run 没有待人工审核的 Claim")
    invalid = db.query(
        "SELECT cl.claim_id,COUNT(ce.evidence_id) AS evidence_count,"
        "SUM(CASE WHEN ev.conflicts=1 THEN 1 ELSE 0 END) AS conflict_count "
        "FROM claims cl LEFT JOIN claim_evidence ce ON ce.claim_id=cl.claim_id "
        "LEFT JOIN evidence ev ON ev.evidence_id=ce.evidence_id "
        "LEFT JOIN companies co ON cl.entity_type='company' AND cl.entity_id=co.company_id "
        "LEFT JOIN contacts ct ON cl.entity_type='contact' AND cl.entity_id=ct.contact_id "
        "LEFT JOIN companies cc ON ct.company_id=cc.company_id "
        "WHERE cl.status='proposed' AND COALESCE(co.run_id,cc.run_id)=? "
        "GROUP BY cl.claim_id HAVING evidence_count=0 OR conflict_count>0",
        (run_id,),
    )
    if invalid:
        raise StateError("批次中存在无证据或有冲突的 Claim，未执行任何批准")
    for index, row in enumerate(claims):
        decide_claim(db, row["claim_id"], "accepted", actor, reason, f"{request_id}-{index + 1}")
    with db.transaction() as connection:
        connection.execute(
            "UPDATE runs SET status='researching',next_action='awaiting_drafts' WHERE run_id=?", (run_id,)
        )
        db.append_event(
            connection,
            event_type="research.claims_approved",
            request_id=request_id,
            run_id=run_id,
            payload={"actor": actor, "count": len(claims)},
        )
    return {"status": "succeeded", "run_id": run_id, "accepted_claims": len(claims), "next_action": "awaiting_drafts"}


def ingest_draft_bundle(db: Database, run_id: str, bundle: dict[str, Any], request_id: str) -> dict[str, Any]:
    with db.transaction():
        return _ingest_draft_bundle(db, run_id, bundle, request_id)


def _ingest_draft_bundle(db: Database, run_id: str, bundle: dict[str, Any], request_id: str) -> dict[str, Any]:
    drafts = bundle.get("drafts")
    if not isinstance(drafts, list) or not drafts:
        raise ValidationError("draft bundle 必须包含 drafts")
    seen_pairs: set[tuple[str, str]] = set()
    for draft in drafts:
        if draft.get("run_id") not in {None, run_id}:
            raise ValidationError("draft bundle 引用了其他 run")
        contact_id = str(draft.get("contact_id", ""))
        channel = str(draft.get("channel", ""))
        pair = (contact_id, channel)
        if not contact_id or channel not in {"email", "linkedin"} or pair in seen_pairs:
            raise ValidationError("同一联系人及渠道只允许一份草稿，按实际入口选择渠道")
        seen_pairs.add(pair)
        if not isinstance(draft.get("outreach_contract", {}), dict):
            raise ValidationError("outreach_contract必须是对象")
        require_one_cta(str(draft.get("body", "")), draft.get("outreach_contract", {}).get("primary_cta", ""))
        if not isinstance(draft.get("evidence_ids"), list) or not draft["evidence_ids"]:
            raise ValidationError("每条草稿必须引用 Evidence")
    created = []
    for index, draft in enumerate(drafts):
        existing = db.query(
            "SELECT * FROM drafts WHERE run_id=? AND contact_id=? AND channel=? ORDER BY created_at DESC",
            (run_id, draft["contact_id"], draft["channel"]),
        )
        if existing:
            current = existing[0]
            changed = (
                current["body"] != str(draft["body"]).strip()
                or current["subject"] != str(draft.get("subject", ""))
                or current["product_version"] != str(draft.get("product_version", ""))
                or json.loads(current["evidence_ids"]) != draft["evidence_ids"]
            )
            stored = db.query("SELECT contract_json FROM draft_contracts WHERE draft_id=?", (current["draft_id"],))
            supplied = draft.get("outreach_contract", {})
            supplied = {**supplied, **{f: draft[f] for f in ("fact_check_status", "voice_check_status") if f in draft}}
            contract = json.loads(stored[0]["contract_json"]) if stored else {}
            qa_fields = {"fact_check_status", "voice_check_status", "fact_checked_by", "voice_checked_by"}
            differences = {k for k, v in supplied.items() if contract.get(k) != v}
            if not changed and stored and current["status"] == "draft" and differences and differences <= qa_fields:
                from .outreach import contract_for

                contact = db.query("SELECT * FROM contacts WHERE contact_id=?", (draft["contact_id"],))[0]
                inputs = {k: v for k, v in contract.items() if k not in {"claim_ids", "evidence_fingerprint"}}
                checked = contract_for(
                    contact, {**draft, "outreach_contract": {**inputs, **supplied}}, contract["claim_ids"]
                )
                checked["evidence_fingerprint"] = contract["evidence_fingerprint"]
                ready = all(checked[f] == "passed" for f in ("fact_check_status", "voice_check_status"))
                with db.transaction() as connection:
                    connection.execute(
                        "UPDATE draft_contracts SET contract_json=? WHERE draft_id=?",
                        (json.dumps(checked, ensure_ascii=False, sort_keys=True), current["draft_id"]),
                    )
                    connection.execute(
                        "UPDATE drafts SET status=?,updated_at=? WHERE draft_id=?",
                        ("review_ready" if ready else "draft", utc_now(), current["draft_id"]),
                    )
                    db.append_event(
                        connection,
                        event_type="draft.qa_reviewed",
                        request_id=request_id,
                        run_id=run_id,
                        entity_type="draft",
                        entity_id=current["draft_id"],
                        payload={f: checked[f] for f in qa_fields},
                    )
                contract = checked
            if changed or not stored or any(contract.get(k) != v for k, v in supplied.items()):
                raise StateError("同一联系人和渠道已有不同草稿，请先人工处理旧草稿")
            created.append(existing[0]["draft_id"])
            continue
        result = create_draft(db, {**draft, "run_id": run_id}, f"{request_id}-{index + 1}")
        created.append(result["draft_id"])
    companies = db.query("SELECT COUNT(*) AS n FROM companies WHERE run_id=?", (run_id,))[0]["n"]
    ready_companies = db.query(
        "SELECT COUNT(DISTINCT company_id) AS n FROM drafts WHERE run_id=? AND status IN ('review_ready','approved','sent')",
        (run_id,),
    )[0]["n"]
    complete = companies > 0 and ready_companies == companies
    with db.transaction() as connection:
        connection.execute(
            "UPDATE runs SET status=?,next_action=? WHERE run_id=?",
            ("reviewing" if complete else "researching", "review_outreach" if complete else "awaiting_drafts", run_id),
        )
        db.append_event(
            connection,
            event_type="research.drafts_ingested",
            request_id=request_id,
            run_id=run_id,
            payload={"count": len(created), "complete": complete},
        )
    return {
        "status": "succeeded" if complete else "needs_review",
        "run_id": run_id,
        "draft_ids": created,
        "next_action": "review_outreach" if complete else "awaiting_drafts",
    }


def research_status(db: Database, run_id: str) -> dict[str, Any]:
    run = db.query("SELECT * FROM runs WHERE run_id=?", (run_id,))
    if not run:
        raise ValidationError("找不到 run")
    companies = db.query("SELECT * FROM companies WHERE run_id=? ORDER BY name_original", (run_id,))
    candidates = db.query("SELECT * FROM candidates WHERE run_id=? ORDER BY discovery_status,name_original", (run_id,))
    stats = {
        "companies": len(companies),
        "contacts": 0,
        "source_runs": db.query("SELECT COUNT(*) AS n FROM source_runs WHERE run_id=?", (run_id,))[0]["n"],
        "evidence": 0,
        "claims": 0,
        "proposed_claims": 0,
        "review_ready_email": db.query(
            "SELECT COUNT(*) AS n FROM drafts WHERE run_id=? AND channel='email' AND status='review_ready'", (run_id,)
        )[0]["n"],
        "review_ready_linkedin": db.query(
            "SELECT COUNT(*) AS n FROM drafts WHERE run_id=? AND channel='linkedin' AND status='review_ready'",
            (run_id,),
        )[0]["n"],
        "research_gaps": 0,
    }
    review = []
    for company in companies:
        cid = company["company_id"]
        contacts = db.query("SELECT * FROM contacts WHERE company_id=? ORDER BY rowid DESC", (cid,))
        evidence = db.query(
            "SELECT * FROM evidence WHERE (entity_type='company' AND entity_id=?) OR (entity_type='contact' AND entity_id IN (SELECT contact_id FROM contacts WHERE company_id=?))",
            (cid, cid),
        )
        claims = db.query(
            "SELECT cl.*,GROUP_CONCAT(ce.evidence_id) AS evidence_ids FROM claims cl LEFT JOIN claim_evidence ce ON ce.claim_id=cl.claim_id "
            "WHERE (cl.entity_type='company' AND cl.entity_id=?) OR (cl.entity_type='contact' AND cl.entity_id IN (SELECT contact_id FROM contacts WHERE company_id=?)) GROUP BY cl.claim_id "
            "ORDER BY CASE cl.status WHEN 'accepted' THEN 0 ELSE 1 END,cl.rowid",
            (cid, cid),
        )
        gaps = db.query(
            "SELECT * FROM research_gaps WHERE (entity_type='company' AND entity_id=?) OR "
            "(entity_type='contact' AND entity_id IN (SELECT contact_id FROM contacts WHERE company_id=?))",
            (cid, cid),
        )
        drafts = db.query(
            "SELECT * FROM drafts WHERE run_id=? AND company_id=? ORDER BY channel,created_at",
            (run_id, cid),
        )
        company_preview = dict(company)
        company_claims = [row for row in claims if row["entity_type"] == "company"]
        for claim in company_claims:
            if claim["status"] in {"proposed", "accepted"}:
                company_preview[claim["field_name"]] = claim["proposed_value"]
        contact_previews = []
        for contact in contacts:
            preview = dict(contact)
            for claim in claims:
                if (
                    claim["entity_type"] == "contact"
                    and claim["entity_id"] == contact["contact_id"]
                    and claim["status"] in {"proposed", "accepted"}
                ):
                    preview[claim["field_name"]] = claim["proposed_value"]
            contact_previews.append(preview)
        claim_statuses = {row["status"] for row in claims}
        decision_state = (
            "待人工批准"
            if "proposed" in claim_statuses
            else "已批准"
            if claim_statuses == {"accepted"}
            else "部分已处理"
            if claim_statuses
            else "无 Claim"
        )
        stats["contacts"] += len(contacts)
        stats["evidence"] += len(evidence)
        stats["claims"] += len(claims)
        stats["proposed_claims"] += sum(row["status"] == "proposed" for row in claims)
        stats["research_gaps"] += len(gaps)
        review.append(
            {
                "company": company,
                "contacts": contacts,
                "company_preview": company_preview,
                "contact_previews": contact_previews,
                "decision_state": decision_state,
                "evidence": evidence,
                "claims": claims,
                "drafts": drafts,
                "replies": db.query(
                    "SELECT r.* FROM replies r JOIN contacts c ON c.contact_id=r.contact_id "
                    "WHERE c.company_id=? ORDER BY julianday(r.received_at) DESC,r.created_at DESC,r.rowid DESC",
                    (cid,),
                ),
                "gaps": gaps,
            }
        )
    return {
        "status": "succeeded",
        "run": run[0],
        "stage": run[0]["next_action"],
        "stats": stats,
        "coverage": discovery_coverage(db, run_id),
        "candidates": candidates,
        "review": review,
    }


def verify_company_list(db: Database, run_id: str, workbook: Path | None = None) -> dict[str, Any]:
    """研究名单的结构/证据门槛；不冒充人工产品判断或外联批准。"""
    snapshot = research_status(db, run_id)
    run = snapshot["run"]
    target = int(run["target_count"])
    failures, warnings = [], list(snapshot["coverage"]["warnings"])
    seen, route_count = set(), 0
    business_route_count = 0
    group_units: dict[str, list[dict[str, Any]]] = {}
    if len(snapshot["review"]) != target:
        failures.append(f"公司数量应为{target}，实际{len(snapshot['review'])}；保持原目标")
    for item in snapshot["review"]:
        company = item["company_preview"]
        name = company["name_original"]
        group = str(company.get("group_key", "")).strip().casefold()
        if group:
            group_units.setdefault(group, []).append(company)
        domain = root_domain(str(company.get("website", "")))
        if not domain or domain in seen:
            failures.append(f"{name}：官网缺失或公司重复")
        seen.add(domain)
        if company["country"] not in str(run["markets"]).split(","):
            failures.append(f"{name}：不在批次市场")
        if not segment_allowed(company["segment"], str(run["segments"])):
            failures.append(f"{name}：不在批次客户类型")
        try:
            validate_company_fit_gates({**company, "name": name})
        except ValidationError as exc:
            failures.append(str(exc))
        if not company["fit_reason"] or company["fit_status"] not in {"priority_chat", "chat_candidate"}:
            failures.append(f"{name}：缺少可解释的业务匹配")
        if not any(
            e["entity_type"] == "company"
            and e["purpose"] == "company_fit"
            and e["access_status"] == "accessible"
            and not e["conflicts"]
            and e["source_level"] in {"primary", "secondary"}
            for e in item["evidence"]
        ):
            failures.append(f"{name}：缺少可访问的公司核验证据")
        if any(e["conflicts"] for e in item["evidence"]):
            failures.append(f"{name}：有来源冲突，需先处理")
        for evidence in item["evidence"]:
            try:
                observed = datetime.fromisoformat(evidence["observed_at"].replace("Z", "+00:00"))
                if observed.tzinfo is None:
                    observed = observed.replace(tzinfo=timezone.utc)
                age = (datetime.now(timezone.utc) - observed).days
                if age < -1:
                    failures.append(f"{name}：来源观察日期在未来")
                elif age > (30 if evidence["entity_type"] == "contact" else 90):
                    warnings.append(f"{name}：旧来源待复核：{evidence['source_url']}")
            except (TypeError, ValueError):
                failures.append(f"{name}：来源观察日期无效")
        if not any(c["selected_company_id"] == company["company_id"] for c in snapshot["candidates"]):
            failures.append(f"{name}：缺少发现候选关联")
        # A published contact page is useful even without a named person or an email.
        route_contacts = {
            e["entity_id"]
            for e in item["evidence"]
            if e["entity_type"] == "contact"
            and e["purpose"] == "contact_route"
            and e["access_status"] == "accessible"
            and not e["conflicts"]
            and str(e["source_url"]).startswith(("https://", "http://"))
        }
        usable_contacts = [
            c
            for c in item["contact_previews"]
            if c["contact_id"] in route_contacts
            and c.get("route_grade", "R3") != "R0"
            and c.get("route_purpose", "unknown") != "unsuitable"
            and c.get("identity_status") != "contradicted"
        ]
        route = bool(usable_contacts)
        business_route_count += int(
            any(c.get("route_grade") in {"R1", "R2"} and c.get("route_purpose") == "business" for c in usable_contacts)
        )
        with db.transaction() as connection:
            for c in item["contact_previews"] or [{}]:
                if active_stop_exists(
                    connection,
                    c.get("contact_id", ""),
                    company["company_id"],
                    preview_email=c.get("public_email", ""),
                    preview_linkedin=c.get("linkedin_url", ""),
                    preview_domain=domain,
                ):
                    failures.append(f"{name}：命中停止联系，不可作为本批可联系推荐")
                    break
        route_count += int(route)
        if not route:
            failures.append(f"{name}：对接入口待补查；保留公司，但名单交付尚未完成")
        if item["gaps"]:
            warnings.extend(f"{name}：{g['business_impact']}" for g in item["gaps"] if g["status"] == "open")
            failures.extend(
                f"{name}：关键公司缺口未解决：{g['business_impact']}"
                for g in item["gaps"]
                if g["status"] == "open" and g["blocking"] and g["entity_type"] == "company"
            )
    for group, members in group_units.items():
        if len(members) > 1 and any(not m.get("independent_unit_reason") for m in members):
            failures.append(f"集团{group}：多个品牌缺独立合作单元依据，不能重复计数")
    if not workbook or not workbook.is_file():
        failures.append("未提供存在的结果Excel")
    else:
        from openpyxl import load_workbook

        from ..exporters.workbook import HEADERS

        wb = load_workbook(workbook, read_only=True)
        try:
            if wb.sheetnames != ["客户结果"] or list(next(wb.active.values)) != HEADERS:
                failures.append("Excel不符合单表7列合同")
            else:
                if any(
                    wb.active.cell(i, 1).value for i in range(len(snapshot["review"]) + 2, len(snapshot["review"]) + 5)
                ):
                    failures.append("Excel公司区含多余记录")
                actual = [wb.active.cell(i, 1).value for i in range(2, len(snapshot["review"]) + 2)]
                expected = [r["company_preview"]["name_original"] for r in snapshot["review"]]
                if sorted(str(v) for v in actual) != sorted(expected):
                    failures.append("Excel公司与本批次不一致")
        finally:
            wb.close()
    return {
        "status": "succeeded" if not failures else "needs_review",
        "passed": not failures,
        "scope": "companies",
        "run_id": run_id,
        "target_count": target,
        "company_count": len(snapshot["review"]),
        "public_route_count": route_count,
        "business_route_count": business_route_count,
        "handoff_route_count": route_count - business_route_count,
        "research_delivery_status": "ready_for_review" if not failures else "needs_research",
        "product_match_status": "requires_business_review" if run["product_version"] else "product_not_provided",
        "outreach_ready": False,
        "proposed_claims": snapshot["stats"]["proposed_claims"],
        "failures": failures,
        "warnings": warnings,
        "boundary": "仅名单结构及证据记录验收；网页语义、产品匹配和发送许可不能由本检查证明",
    }


def verify_research_uat(
    db: Database, run_id: str, workbook: Path | None = None, *, scope: str = "outreach"
) -> dict[str, Any]:
    """按真实纵向闭环验收，而不是只检查表中有没有记录。"""
    if scope == "companies":
        return verify_company_list(db, run_id, workbook)
    if scope != "outreach":
        raise ValidationError("无效验收范围")
    snapshot = research_status(db, run_id)
    run = snapshot["run"]
    stats = snapshot["stats"]
    coverage = snapshot["coverage"]
    failures: list[str] = []
    target = max(1, int(run["target_count"]))
    allowed_markets = {item.strip() for item in str(run["markets"]).split(",") if item.strip()}
    if stats["companies"] < target:
        failures.append(f"公司数不足：{stats['companies']}/{target}")
    if int(run["target_count"]) >= int(run["production_min_target"]) and not coverage["production_ready"]:
        failures.append("正式批次发现覆盖不足；请根据 coverage.warnings 补搜，不要用窄候选池结束任务")
    discovery = coverage["candidate_count"]
    if not discovery or int(discovery) <= stats["companies"]:
        failures.append("没有证明主动发现数量大于最终入选数量")
    manual_runs = db.query("SELECT COUNT(*) AS n FROM source_runs WHERE run_id=? AND adapter_id='manual'", (run_id,))[
        0
    ]["n"]
    if manual_runs:
        failures.append("真实 UAT 不允许依赖 manual adapter")
    source_run_ids = {
        row["source_run_id"] for row in db.query("SELECT source_run_id FROM source_runs WHERE run_id=?", (run_id,))
    }

    for item in snapshot["review"]:
        company = item["company_preview"]
        name = company["name_original"]
        website = str(company["website"] or "")
        host = (urlsplit(website).hostname or "").lower()
        if not host or host.endswith(".test") or "example." in host:
            failures.append(f"{name}：官网不是可接受的真实域名")
        if company["country"] not in allowed_markets:
            failures.append(f"{name}：country 不在本批次市场范围")
        if not segment_allowed(company["segment"], str(run["segments"])):
            failures.append(f"{name}：segment 不在本批次客户类型范围")
        if company["fit_status"] in {"", "research_needed", "not_fit", "do_not_contact"}:
            failures.append(f"{name}：fit_status 未达到外联候选标准")
        if not str(company["fit_reason"] or "").strip():
            failures.append(f"{name}：缺少 fit_reason")
        company_evidence = [row for row in item["evidence"] if row["entity_type"] == "company"]
        if not company_evidence:
            failures.append(f"{name}：缺少公司 Evidence")
        elif any(row["source_run_id"] not in source_run_ids for row in company_evidence):
            failures.append(f"{name}：公司 Evidence 没有归属本 run 的 SourceRun")
        if not any(row["purpose"] == "discovery" for row in company_evidence):
            failures.append(f"{name}：缺少主动发现来源 Evidence")
        if not any(row["purpose"] == "company_fit" for row in company_evidence):
            failures.append(f"{name}：缺少独立公司核验 Evidence")
        if any((urlsplit(str(row["source_url"])).hostname or "").endswith(".test") for row in item["evidence"]):
            failures.append(f"{name}：Evidence 含测试域名")
        contacts = item["contact_previews"]
        if not contacts:
            failures.append(f"{name}：缺少负责人研究结果")
            continue
        contact = contacts[0]
        if contact["identity_status"] not in {"supported", "verified"}:
            failures.append(f"{name}：联系人身份未达到 supported/verified")
        if contact["contact_role"] != "handoff" and (
            not str(contact["full_name"] or "").strip() or not str(contact["title"] or "").strip()
        ):
            failures.append(f"{name}：负责人姓名或职位缺失")
        if not contact["public_email"] and not str(contact["linkedin_url"] or "").startswith(("http://", "https://")):
            failures.append(f"{name}：缺少外联支持的邮箱/LinkedIn入口，官网表单不能替代发送渠道核验")
        contact_evidence = [row for row in item["evidence"] if row["entity_type"] == "contact"]
        if not contact_evidence:
            failures.append(f"{name}：负责人没有身份 Evidence")
        elif any(row["source_run_id"] not in source_run_ids for row in contact_evidence):
            failures.append(f"{name}：联系人 Evidence 没有归属本 run 的 SourceRun")
        accepted = [row for row in item["claims"] if row["status"] == "accepted"]
        if not accepted:
            failures.append(f"{name}：没有经过人工接受的 Claim")
        else:
            accepted_company_fields = {row["field_name"] for row in accepted if row["entity_type"] == "company"}
            required_company_fields = {"website", "country", "segment", "fit_reason", "fit_status"}
            required_company_fields |= (
                {"qualification_policy", *SERVICE_FIELDS}
                if any(
                    row["field_name"] == "qualification_policy" and row["proposed_value"] == "service_fit_v1"
                    for row in accepted
                    if row["entity_type"] == "company"
                )
                else set(_FIT_GATE_VALUES)
            )
            if not required_company_fields.issubset(accepted_company_fields):
                failures.append(f"{name}：公司正式字段没有全部经过 Claim 批准")
            accepted_contact_fields = {
                row["field_name"]
                for row in accepted
                if row["entity_type"] == "contact" and row["entity_id"] == contact["contact_id"]
            }
            required_contact_fields = {"role_scope", "identity_status"}
            if contact["contact_role"] != "handoff":
                required_contact_fields |= {"full_name", "title"}
            if not required_contact_fields.issubset(accepted_contact_fields):
                failures.append(f"{name}：负责人正式字段没有全部经过 Claim 批准")
        drafts = db.query(
            "SELECT draft_id,contact_id,channel,body,evidence_ids,status FROM drafts WHERE run_id=? AND company_id=?",
            (run_id, company["company_id"]),
        )
        ready = {row["channel"]: row for row in drafts if row["status"] == "review_ready"}
        if not ready:
            failures.append(f"{name}：缺少首选渠道可审核草稿")
        else:
            if len(ready) == 2 and ready["email"]["body"].strip() == ready["linkedin"]["body"].strip():
                failures.append(f"{name}：Email 与 LinkedIn 草稿没有体现渠道差异")
            for channel, draft in ready.items():
                try:
                    contracts = db.query(
                        "SELECT contract_json FROM draft_contracts WHERE draft_id=?", (draft["draft_id"],)
                    )
                    contract = json.loads(contracts[0]["contract_json"]) if contracts else {}
                    if any(contract.get(f) != "passed" for f in ("fact_check_status", "voice_check_status")):
                        failures.append(f"{name}：{channel} 事实或语气复核未通过")
                    if not contract.get("fact_checked_by") or not contract.get("voice_checked_by"):
                        failures.append(f"{name}：{channel} 缺少事实或语气复核人")
                    require_one_cta(draft["body"], contract.get("primary_cta", ""))
                except ValidationError:
                    failures.append(f"{name}：{channel} 草稿不是恰好一个 CTA")
                try:
                    evidence_ids = json.loads(draft["evidence_ids"])
                except (TypeError, json.JSONDecodeError):
                    evidence_ids = []
                if not isinstance(evidence_ids, list) or not evidence_ids:
                    failures.append(f"{name}：{channel} 草稿没有 Evidence 引用")
                    continue
                placeholders = ",".join("?" for _ in evidence_ids)
                matched = db.query(
                    f"SELECT COUNT(*) AS n FROM evidence WHERE evidence_id IN ({placeholders}) "
                    "AND ((entity_type='company' AND entity_id=?) OR "
                    "(entity_type='contact' AND entity_id=?))",
                    (*evidence_ids, company["company_id"], draft["contact_id"]),
                )[0]["n"]
                if matched != len(set(evidence_ids)):
                    failures.append(f"{name}：{channel} 草稿 Evidence 不属于该公司或联系人")

    if workbook is not None:
        if not workbook.is_file():
            failures.append("最终 Excel 不存在")
        else:
            try:
                from openpyxl import load_workbook

                opened = load_workbook(workbook, read_only=True, data_only=False)
                if opened.sheetnames != ["客户结果"]:
                    failures.append("Excel 必须只有客户结果一张表")
                elif opened["客户结果"].max_row < stats["companies"] + 1:
                    failures.append("Excel 首页公司行数不足")
                else:
                    headers = {cell.value for cell in opened["客户结果"][1]}
                    required_headers = {
                        "公司",
                        "市场/客群",
                        "为什么值得联系",
                        "找谁",
                        "联系入口",
                        "开发信",
                        "当前结果/下一步",
                    }
                    if not required_headers.issubset(headers):
                        failures.append("Excel 首页总览缺少业务审阅必需列")
                opened.close()
            except ImportError:
                failures.append("缺少 openpyxl，无法验收 Excel")
    else:
        failures.append("未提供最终 Excel 路径")
    return {
        "status": "succeeded" if not failures else "needs_review",
        "passed": not failures,
        "scope": "outreach",
        "run_id": run_id,
        "stage": snapshot["stage"],
        "discovered_count": int(discovery or 0),
        "production_ready": coverage["production_ready"],
        "coverage_warnings": coverage["warnings"],
        "source_performance": coverage["source_performance"],
        "stats": stats,
        "workbook": str(workbook) if workbook else None,
        "failures": failures,
    }


def run_research(
    db: Database,
    run_id: str,
    request_id: str,
    *,
    source_bundle_path: Path | None = None,
    approve_actor: str = "",
    approve_reason: str = "",
    draft_bundle_path: Path | None = None,
) -> dict[str, Any]:
    """同一命令可重复调用；根据当前阶段前进或返回下一步。"""
    with db.transaction() as connection:
        _run_exists(connection, run_id)
        row = connection.execute("SELECT status,next_action FROM runs WHERE run_id=?", (run_id,)).fetchone()
        if row["next_action"] in {"导入或研究候选公司", ""} and not source_bundle_path:
            connection.execute(
                "UPDATE runs SET status='researching',next_action='awaiting_source_execution' WHERE run_id=?", (run_id,)
            )
            db.append_event(connection, event_type="research.started", request_id=request_id, run_id=run_id)
    if source_bundle_path:
        ingest_source_bundle(db, run_id, _read_bundle(source_bundle_path) or {}, request_id)
    if approve_actor:
        approve_run_claims(db, run_id, approve_actor, approve_reason, request_id)
    if draft_bundle_path:
        ingest_draft_bundle(db, run_id, _read_bundle(draft_bundle_path) or {}, request_id)
    result = research_status(db, run_id)
    if result["stage"] == "awaiting_source_execution":
        policy = result["coverage"]["policy"]
        result["executor_request"] = {
            "executor": "codex_bridge",
            "task": "discover_and_inspect",
            "rules": [
                "先建立 source_plan；客户类型分布是补搜提示，不是固定配额",
                f"从互补公开来源主动发现目标公司；本批次来源基线为{policy['min_source_categories']}类",
                f"本批次候选池基线为正式目标的{policy['candidate_pool_multiplier']:g}倍，不使用预写名单",
                "所有候选都要保存；未入选者标记 candidate_pool、held 或 excluded 并说明原因",
                "同一候选的多个发现来源写入 sources，用于计算独有贡献和重叠",
                "必须打开来源页面后保存观察内容",
                "官网优先；LinkedIn 仅使用无需登录的公开信息",
                "不猜邮箱，受限和未找到必须明确记录",
                "新研究统一公司池，qualification_policy=service_fit_v1；不设行业配额，不扩展本批次市场",
                "逐项记录fit_basis、business_signal、service_fit、cooperation_hypothesis、counter_evidence、next_action；假设不伪装为事实",
                "新研究提供逐字段field_sources；contact提供route_grade与route_purpose，普通页面不等于商务负责人",
                "达到数量后做高价值对照搜索再择优；集团多品牌需独立合作依据，不凑数",
                "当地入境DMC、只做本国地接的公司和自营中国地接的一体化竞争者不能冒充潜在买方",
            ],
            "candidate_fields": [
                "name",
                "source_url",
                "observed_text",
                "source_category",
                "query_summary",
                "country_hint",
                "segment_hint",
                "discovery_status",
                "decision_reason",
                "sources",
            ],
        }
    return result
