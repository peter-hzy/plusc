"""Convert the seven-column review workbook into a conservative legacy bundle.

The workbook is an export/review surface, not a database backup.  This module
therefore preserves recoverable text while deliberately creating no Evidence,
Claim, approval, or sent-message records.
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

from ..errors import ValidationError

BUSINESS_SHEET = "客户结果"
HEADERS = ["公司", "市场/客群", "为什么值得联系", "找谁", "联系入口", "开发信", "当前结果/下一步"]
TITLE_SUFFIX = "｜内部文案提案（未发送）"


def _text(value: Any) -> str:
    return str(value or "").strip()


def _stable_id(prefix: str, value: str) -> str:
    return f"{prefix}_{hashlib.sha256(value.encode('utf-8')).hexdigest()[:24]}"


def _first_line(value: Any) -> str:
    return _text(value).splitlines()[0].strip() if _text(value) else ""


def _email_from_cell(cell: Any) -> str:
    target = _text(cell.hyperlink.target if cell.hyperlink else "")
    if target.lower().startswith("mailto:"):
        return target[7:].split("?", 1)[0].strip().lower()
    match = re.search(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", _text(cell.value), re.I)
    return match.group(0).lower() if match else ""


def _public_entry(cell: Any) -> str:
    target = _text(cell.hyperlink.target if cell.hyperlink else "")
    return target if target and not target.startswith("#") else _first_line(cell.value)


def _route_grade(value: Any) -> str:
    text = _text(value)
    explicit = re.search(r"\bR([0-3])\b", text, re.I)
    if explicit:
        return f"R{explicit.group(1)}"
    if "具名职责入口" in text:
        return "R1"
    if "部门职责入口" in text:
        return "R2"
    if "通用转交" in text or "企业转交" in text:
        return "R3"
    return "R0"


def _market(value: Any) -> tuple[str, str]:
    text = _text(value)
    market, _, segment = text.partition("/")
    country = {"澳洲": "AU", "澳大利亚": "AU", "新西兰": "NZ"}.get(market.strip(), market.strip())
    return country, segment.strip()


def _named_contact(value: Any) -> tuple[str, str]:
    text = _text(value)
    if not text or "未找到" in text or "团队" in text or "入口" in text or "请转" in text:
        return "", ""
    name, separator, title = text.partition(" / ")
    return (name.strip(), title.strip()) if separator else ("", "")


def _extract_review_blocks(sheet: Any) -> dict[str, dict[str, str]]:
    blocks: dict[str, dict[str, str]] = {}
    for row in range(2, sheet.max_row + 1):
        title = _text(sheet.cell(row, 1).value)
        if not title.endswith(TITLE_SUFFIX):
            continue
        company = title[: -len(TITLE_SUFFIX)].strip()
        body = _text(sheet.cell(row + 1, 1).value)
        limitation = _text(sheet.cell(row + 2, 1).value)
        audit = _text(sheet.cell(row + 3, 1).value)
        fields: dict[str, str] = {"body": body, "limitation": limitation}
        for line in audit.splitlines():
            key, separator, value = line.partition("：")
            if separator:
                fields[key.strip()] = value.strip()
        blocks[company] = fields
    return blocks


def convert_review_workbook(path: Path, *, run_id: str = "") -> dict[str, Any]:
    """Return a legacy-workbook-1.0 bundle containing only recoverable records."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - covered by doctor/install checks
        raise ValidationError("读取Excel需要安装openpyxl") from exc

    path = path.resolve()
    if not path.is_file():
        raise ValidationError(f"Excel不存在：{path}")
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if BUSINESS_SHEET not in workbook.sheetnames:
            raise ValidationError(f"Excel必须包含工作表：{BUSINESS_SHEET}")
        sheet = workbook[BUSINESS_SHEET]
        headers = [_text(sheet.cell(1, column).value) for column in range(1, 8)]
        if headers != HEADERS:
            raise ValidationError("Excel七列表头与PlusC审阅格式不一致")
        end_row = 1
        for row in range(2, sheet.max_row + 1):
            if not _text(sheet.cell(row, 1).value):
                break
            end_row = row
        if end_row < 2:
            raise ValidationError("Excel没有公司记录")
        blocks = _extract_review_blocks(sheet)
        recovered_run_id = run_id.strip() or f"run_excel_{source_hash[:16]}"
        companies: list[dict[str, Any]] = []
        contacts: list[dict[str, Any]] = []
        gaps: list[dict[str, Any]] = []
        drafts: list[dict[str, Any]] = []
        markets: set[str] = set()
        segments: set[str] = set()
        missing_blocks: list[str] = []
        for row in range(2, end_row + 1):
            name = _text(sheet.cell(row, 1).value)
            website = _text(sheet.cell(row, 1).hyperlink.target if sheet.cell(row, 1).hyperlink else "")
            country, segment = _market(sheet.cell(row, 2).value)
            markets.add(country)
            segments.add(segment)
            fit_reason = _text(sheet.cell(row, 3).value)
            risk_match = re.search(r"(?:限制|风险)：([^\n]+)", fit_reason)
            risk_note = risk_match.group(1).strip() if risk_match else "原始证据、产品能力与合作假设均待人工复核"
            company_id = _stable_id("co", f"{name}|{website}")
            block = blocks.get(name, {})
            if not block:
                missing_blocks.append(name)
            observed_at = block.get("核查日", "")
            companies.append(
                {
                    "company_id": company_id,
                    "run_id": recovered_run_id,
                    "company_name_original": name,
                    "country": country,
                    "company_type": segment,
                    "website": website,
                    "candidate_status": "research_needed",
                    "fit_reason": fit_reason,
                    "risk_note": risk_note,
                    "observed_at": observed_at,
                    "discovery_source_url": block.get("公司依据", ""),
                    "review_state": "pending_verification",
                }
            )
            contact_text = _text(sheet.cell(row, 4).value)
            contact_name, contact_title = _named_contact(contact_text)
            entry_cell = sheet.cell(row, 5)
            email = _email_from_cell(entry_cell)
            public_entry = _public_entry(entry_cell)
            grade = _route_grade(entry_cell.value)
            contact_id = _stable_id("ct", f"{company_id}|{contact_text}|{public_entry}")
            contacts.append(
                {
                    "contact_id": contact_id,
                    "company_id": company_id,
                    "name": contact_name,
                    "title": contact_title,
                    "department_or_scope": f"{contact_text}\n入口等级：{grade}\n公开入口：{public_entry}",
                    "contact_role": "handoff" if grade in {"R2", "R3", "R0"} else "primary",
                    "business_email": email,
                    "identity_status": "unverified",
                    "contact_status": "not_contacted",
                    "observed_at": observed_at,
                }
            )
            for entity_type, entity_id, field_name, impact in (
                ("company", company_id, "original_evidence", "Excel仅保留来源URL和摘要；原始Evidence、Claim与SourceRun未随文件提供，必须重新打开来源核验。"),
                ("company", company_id, "service_fit_approval", "产品能力、合作假设与匹配结论来自同伴审阅快照，尚无ChinaVela人工批准记录。"),
                ("contact", contact_id, "responsible_contact", f"入口等级暂按Excel文字解析为{grade}；身份、职责、邮箱归属和商务用途均待重新核验。"),
            ):
                gaps.append(
                    {
                        "gap_id": _stable_id("gap", f"{entity_id}|{field_name}"),
                        "entity_type": entity_type,
                        "entity_id": entity_id,
                        "missing_field": field_name,
                        "why_it_matters": impact,
                        "blocks_outreach": True,
                        "gap_status": "open",
                    }
                )
            body = block.get("body", "")
            if body:
                subject = ""
                match = re.match(r"Subject:\s*([^\n]+)\s*\n+", body, re.I)
                if match:
                    subject = match.group(1).strip()
                    body = body[match.end() :].strip()
                drafts.append(
                    {
                        "message_id": _stable_id("msg", f"{contact_id}|{subject}|{body}"),
                        "contact_id": contact_id,
                        "channel": "email",
                        "subject": subject,
                        "message": body,
                        "draft_status": "draft",
                        "human_review_status": "pending",
                        "fact_check_status": "needs_review",
                        "voice_check_status": "needs_review",
                        "source_note": "Recovered from review workbook; not approved and not sent.",
                    }
                )
        return {
            "schema_version": "1.0",
            "conversion": {
                "source_path": str(path),
                "source_sha256": source_hash,
                "source_format": "plusc-seven-column-review-workbook",
                "company_count": len(companies),
                "draft_count": len(drafts),
                "detected_segments": sorted(segments),
                "missing_review_blocks": missing_blocks,
                "limitations": [
                    "No Evidence, Claim, SourceRun, approval, reply, or sent state was reconstructed.",
                    "All company, contact, product-fit, and draft facts remain pending human verification.",
                ],
            },
            "sheets": {
                "RunConfig": [
                    {
                        "run_id": recovered_run_id,
                        "created_at": "",
                        "market": ",".join(sorted(markets)),
                        "client_types": "*",
                        "product_profile_version": "",
                        "voice_profile_version": "",
                        "target_count": len(companies),
                        "actual_company_count": len(companies),
                        "operator": "ChinaVela workbook takeover",
                        "run_status": "review_ready",
                    }
                ],
                "Companies": companies,
                "Contacts": contacts,
                "Evidence": [],
                "ResearchGaps": gaps,
                "Drafts": drafts,
                "Replies": [],
                "StopList": [],
                "DuplicateReview": [],
            },
        }
    finally:
        workbook.close()
