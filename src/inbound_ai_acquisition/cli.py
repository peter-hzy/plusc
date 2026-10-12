"""`inbound-leads` 统一命令行入口。"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import sqlite3
import sys
import uuid
from pathlib import Path
from typing import Any

from . import __version__
from .adapters.registry import built_in_adapters
from .application.discovery import discovery_coverage
from .application.legacy_import import import_legacy_bundle
from .application.research import research_status, run_research, verify_research_uat
from .application.services import (
    add_stop,
    approve_draft,
    backup_database,
    close_run,
    create_draft,
    create_run,
    decide_claim,
    import_companies,
    import_contacts,
    import_evidence,
    ingest_reply,
    record_source_result,
    release_stop,
    run_status,
    scan_duplicates,
    set_draft_status,
)
from .audit.public_repo import audit_public_repo
from .audit.release import audit_release, build_lock
from .config import Workspace
from .errors import AcquisitionError, ValidationError
from .exporters.json_export import export_json
from .exporters.workbook import export_workbook
from .storage.database import Database, canonical_hash, new_id, utc_now

PROJECT_ROOT = Path(__file__).resolve().parents[2]


def configure_console() -> None:
    """统一命令行输出编码，避免 Windows 默认代码页破坏中文 JSON。"""
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="backslashreplace")


def emit(payload: Any) -> None:
    print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))


def request_id(args: argparse.Namespace) -> str:
    return getattr(args, "request_id", None) or f"req_{uuid.uuid4().hex[:16]}"


def workspace_db(args: argparse.Namespace, require: bool = True) -> tuple[Workspace, Database]:
    root = Path(args.workspace).expanduser().resolve()
    workspace = Workspace(root)
    db = Database(workspace.database)
    if require and db.schema_version() is None:
        raise ValidationError(f"工作区未初始化：{root}")
    return workspace, db


def add_workspace(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--workspace", required=True, help="私有工作区路径")


def add_request(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--request-id")


def add_actor_reason(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--actor", required=True)
    parser.add_argument("--reason", required=True)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="inbound-leads", description="入境游 B2B AI 获客本地工作流")
    parser.add_argument("--version", action="version", version=__version__)
    commands = parser.add_subparsers(dest="command", required=True)

    doctor = commands.add_parser("doctor")
    doctor.add_argument("--workspace")
    init = commands.add_parser("init")
    add_workspace(init)
    migrate = commands.add_parser("migrate")
    add_workspace(migrate)
    backup = commands.add_parser("backup")
    add_workspace(backup)
    backup.add_argument("--output")

    context = commands.add_parser("context").add_subparsers(dest="context_command", required=True)
    context_validate = context.add_parser("validate")
    add_workspace(context_validate)

    run = commands.add_parser("run").add_subparsers(dest="run_command", required=True)
    run_create = run.add_parser("create")
    add_workspace(run_create)
    add_request(run_create)
    run_create.add_argument("--market", required=True)
    run_create.add_argument("--segments", default="*", help="默认不限行业；仅用户明确限制时指定类型")
    run_create.add_argument("--target", type=int, default=20)
    run_create.add_argument("--research-only", action="store_true", help="产品未提供时仅研究公司，不伪造产品或准备发送")
    run_create.add_argument("--candidate-pool-multiplier", type=float, default=3.0)
    run_create.add_argument("--min-source-categories", type=int, default=3)
    run_create.add_argument("--production-min-target", type=int, default=20)
    run_create.add_argument("--idempotency-key")
    run_status_parser = run.add_parser("status")
    add_workspace(run_status_parser)
    run_status_parser.add_argument("--run", required=True)
    run_close = run.add_parser("close")
    add_workspace(run_close)
    add_request(run_close)
    run_close.add_argument("--run", required=True)

    research = commands.add_parser("research").add_subparsers(dest="research_command", required=True)
    research_run = research.add_parser("run")
    add_workspace(research_run)
    add_request(research_run)
    research_run.add_argument("--run", required=True)
    research_run.add_argument("--source-bundle", help="Codex bridge 内部研究结果；普通用户无需手工制作")
    research_run.add_argument("--approve-actor", help="真实人工审核人；仅在批次确认后填写")
    research_run.add_argument("--approve-reason", default="")
    research_run.add_argument("--draft-bundle", help="Codex bridge 内部草稿结果；普通用户无需手工制作")
    research_status_parser = research.add_parser("status")
    add_workspace(research_status_parser)
    research_status_parser.add_argument("--run", required=True)
    research_coverage = research.add_parser("coverage")
    add_workspace(research_coverage)
    research_coverage.add_argument("--run", required=True)
    research_verify = research.add_parser("verify")
    add_workspace(research_verify)
    research_verify.add_argument("--run", required=True)
    research_verify.add_argument("--workbook")
    research_verify.add_argument("--scope", choices=("companies", "outreach"), default="outreach")

    imports = commands.add_parser("import").add_subparsers(dest="import_command", required=True)
    for name in ("companies", "contacts", "evidence", "legacy", "review"):
        item = imports.add_parser(name)
        add_workspace(item)
        add_request(item)
        item.add_argument("--file", required=True)
        if name == "companies":
            item.add_argument("--run", required=True)

    source = commands.add_parser("source").add_subparsers(dest="source_command", required=True)
    for name in ("list", "status", "record", "execute"):
        item = source.add_parser(name)
        add_workspace(item)
        if name == "execute":
            item.add_argument("--adapter", choices=tuple(built_in_adapters()), required=True)
            item.add_argument("--operation", choices=("discover", "inspect_company", "inspect_contact"), required=True)
            item.add_argument("--file", required=True)
        if name == "status":
            item.add_argument("--run")
        if name == "record":
            item.add_argument("--file", required=True)
            item.add_argument("--idempotency-key", required=True)
            add_request(item)

    for noun in ("company", "contact"):
        group = commands.add_parser(noun).add_subparsers(dest=f"{noun}_command", required=True)
        listing = group.add_parser("list")
        add_workspace(listing)
        listing.add_argument("--run")
        show = group.add_parser("show")
        add_workspace(show)
        show.add_argument("--id", required=True)

    claim = commands.add_parser("claim").add_subparsers(dest="claim_command", required=True)
    claim_list = claim.add_parser("list")
    add_workspace(claim_list)
    claim_list.add_argument("--status")
    propose = claim.add_parser("propose")
    add_workspace(propose)
    add_request(propose)
    propose.add_argument("--entity-type", choices=("company", "contact"), required=True)
    propose.add_argument("--entity-id", required=True)
    propose.add_argument("--field", required=True)
    propose.add_argument("--value", required=True)
    propose.add_argument("--evidence", action="append", default=[])
    propose.add_argument("--grade", default="unknown")
    for name in ("accept", "hold", "reject"):
        item = claim.add_parser(name)
        add_workspace(item)
        add_request(item)
        item.add_argument("--claim", required=True)
        add_actor_reason(item)

    dedupe = commands.add_parser("dedupe").add_subparsers(dest="dedupe_command", required=True)
    dedupe_scan = dedupe.add_parser("scan")
    add_workspace(dedupe_scan)
    add_request(dedupe_scan)
    dedupe_review = dedupe.add_parser("review")
    add_workspace(dedupe_review)
    add_request(dedupe_review)
    dedupe_review.add_argument("--candidate", required=True)
    dedupe_review.add_argument("--decision", choices=("same", "different"), required=True)
    add_actor_reason(dedupe_review)

    draft = commands.add_parser("draft").add_subparsers(dest="draft_command", required=True)
    draft_create = draft.add_parser("create")
    add_workspace(draft_create)
    add_request(draft_create)
    draft_create.add_argument("--file", required=True)
    draft_check = draft.add_parser("check")
    add_workspace(draft_check)
    draft_check.add_argument("--draft", required=True)
    draft_approve = draft.add_parser("approve")
    add_workspace(draft_approve)
    add_request(draft_approve)
    draft_approve.add_argument("--draft", required=True)
    add_actor_reason(draft_approve)
    for name in ("mark-sent", "stop"):
        item = draft.add_parser(name)
        add_workspace(item)
        add_request(item)
        item.add_argument("--draft", required=True)

    reply = commands.add_parser("reply").add_subparsers(dest="reply_command", required=True)
    reply_ingest = reply.add_parser("ingest")
    add_workspace(reply_ingest)
    add_request(reply_ingest)
    reply_ingest.add_argument("--file", required=True)
    reply_ingest.add_argument("--idempotency-key", required=True)

    stop = commands.add_parser("stop").add_subparsers(dest="stop_command", required=True)
    stop_add = stop.add_parser("add")
    add_workspace(stop_add)
    add_request(stop_add)
    stop_add.add_argument("--scope-type", required=True)
    stop_add.add_argument("--scope-value", required=True)
    add_actor_reason(stop_add)
    stop_list = stop.add_parser("list")
    add_workspace(stop_list)
    stop_list.add_argument("--all", action="store_true")
    stop_release = stop.add_parser("release")
    add_workspace(stop_release)
    add_request(stop_release)
    stop_release.add_argument("--rule", required=True)
    add_actor_reason(stop_release)

    export = commands.add_parser("export").add_subparsers(dest="export_command", required=True)
    for name in ("workbook", "json"):
        item = export.add_parser(name)
        add_workspace(item)
        item.add_argument("--run")
        item.add_argument("--output", required=True)

    audit = commands.add_parser("audit").add_subparsers(dest="audit_command", required=True)
    audit_workspace = audit.add_parser("workspace")
    add_workspace(audit_workspace)
    audit_public = audit.add_parser("public-repo")
    audit_public.add_argument("--root", default=str(PROJECT_ROOT))
    audit_lock = audit.add_parser("build-lock")
    audit_lock.add_argument("--root", default=str(PROJECT_ROOT))
    audit_lock.add_argument("--output")
    audit_release_parser = audit.add_parser("release")
    audit_release_parser.add_argument("--root", default=str(PROJECT_ROOT))
    audit_release_parser.add_argument("--lock")

    skill = commands.add_parser("skill").add_subparsers(dest="skill_command", required=True)
    skill_install = skill.add_parser("install")
    skill_install.add_argument("--target", choices=("codex",), default="codex")
    skill_install.add_argument("--replace", action="store_true")
    skill_install.add_argument("--path", help="覆盖默认 ~/.codex/skills")
    skill_status = skill.add_parser("status")
    skill_status.add_argument("--path")
    return parser


def _context_validate(workspace: Workspace, db: Database, *, product_only: bool = False) -> dict[str, Any]:
    profiles = ("product-profile.md",) if product_only else ("product-profile.md", "voice-profile.md")
    required = ("config.toml", *profiles)
    missing = [name for name in required if not (workspace.root / name).exists()]
    placeholders = []
    profile_gaps: dict[str, list[str]] = {}
    required_profile_fields = {
        "product-profile.md": ("版本：", "- 可接待目的地：", "- 可提供服务：", "- 语言："),
        "voice-profile.md": ("版本：", "- 希望像谁在说话："),
    }
    for name in profiles:
        path = workspace.root / name
        if path.exists():
            content = path.read_text(encoding="utf-8")
            if "待填写" in content or "TODO" in content:
                placeholders.append(name)
            lines = [line.strip() for line in content.splitlines()]
            gaps = []
            for label in required_profile_fields[name]:
                matching = [line for line in lines if line.startswith(label)]
                if not matching or not matching[0][len(label) :].strip():
                    gaps.append(label.rstrip("："))
            if gaps:
                profile_gaps[name] = gaps
    return {
        "status": "succeeded" if not missing and not placeholders and not profile_gaps else "needs_review",
        "missing": missing,
        "placeholders": placeholders,
        "profile_gaps": profile_gaps,
        "schema_version": db.schema_version(),
    }


def _profile_version(path: Path) -> str:
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip().startswith("版本："):
            return line.strip().removeprefix("版本：").strip()
    return ""


def _init_workspace(path: Path) -> dict[str, Any]:
    workspace = Workspace(path.resolve())
    created = workspace.ensure_layout()
    for name, source_name in (
        ("product-profile.md", "product-profile.template.md"),
        ("voice-profile.md", "voice-profile.template.md"),
    ):
        target = workspace.root / name
        if not target.exists():
            shutil.copyfile(PROJECT_ROOT / "templates" / source_name, target)
            created.append(name)
    db_result = Database(workspace.database).migrate()
    if "acquisition.db" not in created:
        created.append("acquisition.db")
    return {"status": "succeeded", "workspace": str(workspace.root), "created": created, **db_result}


def _doctor(workspace_path: str | None) -> dict[str, Any]:
    try:
        import openpyxl  # noqa: F401

        excel = True
    except ImportError:
        excel = False
    result: dict[str, Any] = {
        "status": "succeeded",
        "project_version": __version__,
        "python": sys.version.split()[0],
        "python_supported": sys.version_info >= (3, 11),
        "excel_available": excel,
        "optional_api_keys": {
            key: "configured" if os.environ.get(key) else "not_configured"
            for key in ("APOLLO_API_KEY", "HUNTER_API_KEY", "BRAVE_SEARCH_API_KEY")
        },
    }
    if workspace_path:
        workspace = Workspace(Path(workspace_path).resolve())
        result["workspace"] = str(workspace.root)
        result["workspace_exists"] = workspace.root.exists()
        result["workspace_writable"] = os.access(workspace.root, os.W_OK) if workspace.root.exists() else False
        result["schema_version"] = Database(workspace.database).schema_version()
    return result


def _propose_claim(db: Database, args: argparse.Namespace, rid: str) -> dict[str, Any]:
    with db.transaction() as connection:
        table, column = ("companies", "company_id") if args.entity_type == "company" else ("contacts", "contact_id")
        if not connection.execute(f"SELECT 1 FROM {table} WHERE {column}=?", (args.entity_id,)).fetchone():
            raise ValidationError("Claim 引用不存在的实体")
        claim_id = new_id("cl")
        connection.execute(
            "INSERT INTO claims VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                claim_id,
                args.entity_type,
                args.entity_id,
                args.field,
                args.value,
                "proposed",
                args.grade,
                "",
                None,
                None,
                "",
                None,
                utc_now(),
            ),
        )
        for evidence_id in args.evidence:
            if not connection.execute("SELECT 1 FROM evidence WHERE evidence_id=?", (evidence_id,)).fetchone():
                raise ValidationError(f"找不到 evidence：{evidence_id}")
            connection.execute("INSERT INTO claim_evidence VALUES (?,?)", (claim_id, evidence_id))
        db.append_event(
            connection, event_type="claim.proposed", request_id=rid, entity_type="claim", entity_id=claim_id
        )
    return {
        "request_id": rid,
        "status": "succeeded",
        "claim_id": claim_id,
        "changed_records": 1,
        "next_action": "人工接受、暂缓或拒绝",
        "errors": [],
    }


def _dedupe_review(db: Database, args: argparse.Namespace, rid: str) -> dict[str, Any]:
    with db.transaction() as connection:
        cursor = connection.execute(
            "UPDATE duplicate_candidates SET status=?,reviewed_by=?,reviewed_at=?,reason=? WHERE duplicate_id=? AND status='pending'",
            (args.decision, args.actor, utc_now(), args.reason, args.candidate),
        )
        if cursor.rowcount != 1:
            raise ValidationError("找不到待审核重复候选")
        db.append_event(
            connection,
            event_type="dedupe.reviewed",
            request_id=rid,
            entity_type="duplicate_candidate",
            entity_id=args.candidate,
            payload={"decision": args.decision},
        )
    return {"request_id": rid, "status": "succeeded", "changed_records": 1, "errors": []}


def _skill_paths(path_override: str | None) -> tuple[Path, Path]:
    source = PROJECT_ROOT
    base = Path(path_override).expanduser().resolve() if path_override else Path.home() / ".codex" / "skills"
    return source, base / "inbound-ai-acquisition"


def _skill_manifest(root: Path) -> dict[str, str]:
    files: dict[str, str] = {}
    for relative in [Path("SKILL.md"), Path("agents/openai.yaml")]:
        path = root / relative
        if path.exists():
            files[relative.as_posix()] = canonical_hash(path.read_text(encoding="utf-8"))
    for folder in ("references", "templates"):
        for path in sorted((root / folder).glob("*")):
            if path.is_file():
                files[path.relative_to(root).as_posix()] = canonical_hash(path.read_text(encoding="utf-8"))
    return files


def _skill_status(path_override: str | None) -> dict[str, Any]:
    source, target = _skill_paths(path_override)
    if not target.exists():
        return {"status": "not_installed", "target": str(target), "in_sync": False}
    source_manifest, target_manifest = _skill_manifest(source), _skill_manifest(target)
    changed = sorted(
        key
        for key in set(source_manifest) | set(target_manifest)
        if source_manifest.get(key) != target_manifest.get(key)
    )
    unexpected = sorted(
        p.relative_to(target).as_posix()
        for p in target.rglob("*")
        if p.is_file() and "__pycache__" not in p.parts and p.relative_to(target).as_posix() not in source_manifest
    )
    return {
        "status": "succeeded",
        "target": str(target),
        "in_sync": not changed and not unexpected,
        "changed_files": changed,
        "unexpected_files": unexpected,
    }


def _skill_install(path_override: str | None, replace: bool) -> dict[str, Any]:
    source, target = _skill_paths(path_override)
    before = _skill_status(path_override)
    if target.exists() and not replace and not before.get("in_sync"):
        return {**before, "status": "needs_review", "message": "安装版不同；加 --replace 才会覆盖"}
    if before.get("in_sync"):
        return {**before, "backup": None}
    if target.is_symlink() or target.resolve().parent != target.parent.resolve():
        raise ValidationError("安装目标不是预期的独立Skill目录")
    # 先完整构建暂存包；复制失败时现有安装不变。
    target.parent.mkdir(parents=True, exist_ok=True)
    staging = target.parent / f".inbound-acquisition-stage-{uuid.uuid4().hex[:12]}"
    staging.mkdir()
    for relative in (Path("SKILL.md"), Path("agents/openai.yaml")):
        destination = staging / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source / relative, destination)
    for folder in ("references", "templates"):
        shutil.copytree(source / folder, staging / folder)
    backup = None
    if target.exists() and replace and not before.get("in_sync"):
        # 整包移至同级不可发现的备份目录，保留未知文件，避免旧脚本混入新入口。
        backup_root = target.parent / ".inbound-acquisition-backups"
        backup_root.mkdir(exist_ok=True)
        backup = backup_root / f"{target.name}-{uuid.uuid4().hex[:12]}"
        if target.is_symlink() or target.resolve().parent != target.parent.resolve():
            raise ValidationError("安装目标不是预期的独立Skill目录")
        target.rename(backup)
    try:
        staging.rename(target)
    except OSError:
        if backup and not target.exists():
            backup.rename(target)
        raise
    return {
        "status": "succeeded",
        "target": str(target),
        "backup": str(backup) if backup else None,
        "in_sync": _skill_status(path_override)["in_sync"],
    }


def dispatch(args: argparse.Namespace) -> Any:
    command = args.command
    if command == "doctor":
        return _doctor(args.workspace)
    if command == "init":
        return _init_workspace(Path(args.workspace))
    if command == "audit":
        root = Path(getattr(args, "root", PROJECT_ROOT)).resolve()
        if args.audit_command == "public-repo":
            return audit_public_repo(root)
        if args.audit_command == "build-lock":
            return build_lock(root, Path(args.output).resolve() if args.output else root / "skills.lock.json")
        if args.audit_command == "release":
            return audit_release(root, Path(args.lock).resolve() if args.lock else root / "skills.lock.json")
        _, db = workspace_db(args)
        return {
            "status": "succeeded",
            "integrity": db.query("PRAGMA integrity_check"),
            "schema_version": db.schema_version(),
        }
    if command == "skill":
        return _skill_install(args.path, args.replace) if args.skill_command == "install" else _skill_status(args.path)

    workspace, db = workspace_db(args, require=command != "migrate")
    rid = request_id(args)
    if command == "migrate":
        return db.migrate()
    if command == "backup":
        output = (
            Path(args.output).resolve()
            if args.output
            else workspace.root / "backups" / f"acquisition-{utc_now()[:10]}.db"
        )
        return backup_database(db, output)
    if command == "context":
        return _context_validate(workspace, db)
    if command == "run":
        if args.run_command == "create":
            context_status = _context_validate(workspace, db, product_only=True)
            if context_status["status"] != "succeeded" and not args.research_only:
                raise ValidationError(f"产品资料尚未完成：{context_status['profile_gaps']}")
            return create_run(
                db,
                markets=[x.strip() for x in args.market.split(",") if x.strip()],
                segments=[x.strip() for x in args.segments.split(",") if x.strip()],
                target=args.target,
                request_id=rid,
                idempotency_key=args.idempotency_key or rid,
                product_version="" if args.research_only else _profile_version(workspace.root / "product-profile.md"),
                voice_version="" if args.research_only else _profile_version(workspace.root / "voice-profile.md"),
                candidate_pool_multiplier=args.candidate_pool_multiplier,
                min_source_categories=args.min_source_categories,
                production_min_target=args.production_min_target,
            )
        if args.run_command == "status":
            return run_status(db, args.run)
        return close_run(db, args.run, rid)
    if command == "research":
        if args.research_command == "status":
            return research_status(db, args.run)
        if args.research_command == "coverage":
            return discovery_coverage(db, args.run)
        if args.research_command == "verify":
            return verify_research_uat(
                db,
                args.run,
                Path(args.workbook).resolve() if args.workbook else None,
                scope=args.scope,
            )
        return run_research(
            db,
            args.run,
            rid,
            source_bundle_path=Path(args.source_bundle).resolve() if args.source_bundle else None,
            approve_actor=args.approve_actor or "",
            approve_reason=args.approve_reason,
            draft_bundle_path=Path(args.draft_bundle).resolve() if args.draft_bundle else None,
        )
    if command == "import":
        path = Path(args.file).resolve()
        if args.import_command == "companies":
            return import_companies(db, args.run, path, rid)
        if args.import_command == "contacts":
            return import_contacts(db, path, rid)
        if args.import_command == "evidence":
            return import_evidence(db, path, rid)
        if args.import_command == "legacy":
            return import_legacy_bundle(db, path, rid)
        return {
            "status": "needs_review",
            "preview": json.loads(path.read_text(encoding="utf-8")),
            "message": "V1 review 导入只生成预览，不自动覆盖数据库",
        }
    if command == "source":
        if args.source_command == "list":
            adapters = built_in_adapters()
            return {
                "status": "succeeded",
                "adapters": sorted(adapters),
                "capabilities": [
                    a.status()
                    if hasattr(a, "status")
                    else {"adapter_id": a.adapter_id, "capabilities": sorted(a.capabilities)}
                    for a in adapters.values()
                ],
            }
        if args.source_command == "execute":
            from dataclasses import asdict

            adapter = built_in_adapters()[args.adapter]
            request = json.loads(Path(args.file).read_text(encoding="utf-8"))
            if not isinstance(request, dict):
                raise ValidationError("来源请求必须是JSON对象")
            return asdict(getattr(adapter, args.operation)(request))
        if args.source_command == "record":
            return record_source_result(
                db, json.loads(Path(args.file).read_text(encoding="utf-8")), args.idempotency_key, rid
            )
        query = "SELECT * FROM source_runs" + (" WHERE run_id=?" if args.run else "")
        return {"status": "succeeded", "source_runs": db.query(query, (args.run,) if args.run else ())}
    if command in {"company", "contact"}:
        table, id_column = ("companies", "company_id") if command == "company" else ("contacts", "contact_id")
        subcommand = getattr(args, f"{command}_command")
        if subcommand == "show":
            rows = db.query(f"SELECT * FROM {table} WHERE {id_column}=?", (args.id,))
            return {"status": "succeeded" if rows else "not_found", command: rows[0] if rows else None}
        if command == "company" and args.run:
            rows = db.query("SELECT * FROM companies WHERE run_id=?", (args.run,))
        elif command == "contact" and args.run:
            rows = db.query(
                "SELECT c.* FROM contacts c JOIN companies x ON x.company_id=c.company_id WHERE x.run_id=?", (args.run,)
            )
        else:
            rows = db.query(f"SELECT * FROM {table}")
        return {"status": "succeeded", "count": len(rows), f"{command}s": rows}
    if command == "claim":
        if args.claim_command == "list":
            rows = db.query(
                "SELECT * FROM claims" + (" WHERE status=?" if args.status else ""),
                (args.status,) if args.status else (),
            )
            return {"status": "succeeded", "claims": rows}
        if args.claim_command == "propose":
            return _propose_claim(db, args, rid)
        return decide_claim(
            db,
            args.claim,
            {"accept": "accepted", "hold": "held", "reject": "rejected"}[args.claim_command],
            args.actor,
            args.reason,
            rid,
        )
    if command == "dedupe":
        return scan_duplicates(db, rid) if args.dedupe_command == "scan" else _dedupe_review(db, args, rid)
    if command == "draft":
        if args.draft_command == "create":
            return create_draft(db, json.loads(Path(args.file).read_text(encoding="utf-8")), rid)
        if args.draft_command == "check":
            rows = db.query("SELECT * FROM drafts WHERE draft_id=?", (args.draft,))
            return {"status": "succeeded" if rows else "not_found", "draft": rows[0] if rows else None}
        if args.draft_command == "approve":
            return approve_draft(db, args.draft, args.actor, args.reason, rid)
        return set_draft_status(db, args.draft, "sent" if args.draft_command == "mark-sent" else "stopped", rid)
    if command == "reply":
        return ingest_reply(db, json.loads(Path(args.file).read_text(encoding="utf-8")), args.idempotency_key, rid)
    if command == "stop":
        if args.stop_command == "add":
            return add_stop(db, args.scope_type, args.scope_value, args.reason, args.actor, rid)
        if args.stop_command == "release":
            return release_stop(db, args.rule, args.actor, args.reason, rid)
        return {
            "status": "succeeded",
            "stop_rules": db.query("SELECT * FROM stop_rules" + ("" if args.all else " WHERE status='active'")),
        }
    if command == "export":
        return (
            export_workbook(db, Path(args.output).resolve(), args.run)
            if args.export_command == "workbook"
            else export_json(db, Path(args.output).resolve(), args.run)
        )
    raise ValidationError("未知命令")


def main(argv: list[str] | None = None) -> int:
    configure_console()
    try:
        emit(dispatch(build_parser().parse_args(argv)))
        return 0
    except (AcquisitionError, FileNotFoundError, json.JSONDecodeError, sqlite3.Error) as exc:
        emit({"status": "failed", "errors": [str(exc)]})
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
