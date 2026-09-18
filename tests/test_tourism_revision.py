"""隔离测试数据；覆盖旅游客户工作流的真实失败边界。"""

import json
import os
import unittest
from io import BytesIO
from unittest.mock import patch
from urllib.error import HTTPError

import test_cli_workflow as fixtures

from inbound_ai_acquisition.adapters.api import ApiAdapter
from inbound_ai_acquisition.application.research import ingest_draft_bundle, validate_company_fit_gates
from inbound_ai_acquisition.application.services import approve_draft, create_draft, create_run, set_draft_status
from inbound_ai_acquisition.cli import _skill_install, _skill_status
from inbound_ai_acquisition.errors import StateError, ValidationError
from inbound_ai_acquisition.exporters.json_export import export_json
from inbound_ai_acquisition.exporters.workbook import _route_kind, export_workbook


class TourismRevisionTests(unittest.TestCase):
    setUp = fixtures.CliWorkflowTests.setUp
    tearDown = fixtures.CliWorkflowTests.tearDown
    write_json = fixtures.CliWorkflowTests.write_json
    seed_company_contact_claim = fixtures.CliWorkflowTests.seed_company_contact_claim

    def test_incomplete_batch_keeps_target_and_reports_shortfall(self):
        from inbound_ai_acquisition.application.discovery import discovery_coverage

        run = create_run(
            self.db,
            markets=["AU", "NZ"],
            segments=["tour_operator"],
            target=20,
            request_id="fixed-target",
            idempotency_key="fixed-target",
        )["run_id"]
        result = discovery_coverage(self.db, run)
        self.assertEqual(result["target_count"], 20)
        self.assertEqual(result["selected_shortfall"], 20)
        self.assertFalse(result["selection_target_met"])
        self.assertFalse(result["production_ready"])

    def draft_payload(self):
        return {
            "run_id": self.run,
            "contact_id": "ct_demo",
            "channel": "email",
            "subject": "China support",
            "body": "I saw your China tours. Would a regional outline help?",
            "evidence_ids": ["ev_demo"],
            "product_version": "test-v1",
            "outreach_contract": {
                "recipient_kind": "person",
                "address_status": "published",
                "role_reason": "China product owner",
                "fact_checked_by": "test-human",
                "product_approved_by": "test-human",
                "consent": {
                    "basis": "express",
                    "market": "AU",
                    "source_reference": "test-only",
                    "reason": "isolated fixture",
                    "reviewed_by": "test-human",
                    "valid_until": "2099-01-01T00:00:00+00:00",
                },
            },
        }

    def test_directory_host_does_not_merge_distinct_companies(self):
        from inbound_ai_acquisition.application.research import ingest_source_bundle

        rows = [
            {
                "name": name,
                "source_url": "https://directory.example.org/" + name,
                "source_category": "industry_directory",
                "observed_text": "Test directory entry",
            }
            for name in ("Alpha", "Beta")
        ]
        bundle = {
            "executor": "codex_bridge",
            "discovered_count": 2,
            "discovery_query": "fixture",
            "discovered_candidates": rows,
            "companies": [],
        }
        ingest_source_bundle(self.db, self.run, bundle, "directory")
        self.assertEqual(len(self.db.query("SELECT * FROM candidates")), 2)
        rows[0]["website"] = "https://alpha.example.org"
        ingest_source_bundle(self.db, self.run, bundle, "directory-repeat")
        self.assertEqual(len(self.db.query("SELECT * FROM candidates")), 2)
        self.assertEqual(
            self.db.query("SELECT root_domain FROM candidates WHERE name_original='Alpha'")[0]["root_domain"],
            "alpha.example.org",
        )

    def test_workbook_preserves_named_person_and_verified_contact_page(self):
        from openpyxl import load_workbook

        self.seed_company_contact_claim()
        output = self.root / "routes.xlsx"
        export_workbook(self.db, output, self.run)
        wb = load_workbook(output)
        self.assertIn("Alex", wb.active["D2"].value)
        wb.close()
        with self.db.transaction() as c:
            c.execute(
                "UPDATE contacts SET contact_role='handoff',public_email='',linkedin_url='',role_scope='建议转中国产品负责人'"
            )
            c.execute(
                "UPDATE evidence SET entity_type='contact',entity_id='ct_demo',purpose='contact_route',"
                "source_url='https://example.test/contact'"
            )
        export_workbook(self.db, output, self.run)
        wb = load_workbook(output)
        self.assertIn("建议转中国产品负责人", wb.active["D2"].value)
        self.assertIn("官网联系页", wb.active["E2"].value)
        self.assertEqual(wb.active["E2"].hyperlink.target, "https://example.test/contact")
        wb.close()

    def test_subject_and_contract_changes_invalidate_approval(self):
        self.seed_company_contact_claim()
        draft = create_draft(self.db, self.draft_payload(), "new")["draft_id"]
        approve_draft(self.db, draft, "human", "test", "approve")
        with self.db.transaction() as c:
            c.execute("UPDATE drafts SET subject=? WHERE draft_id=?", ("changed subject", draft))
        with self.assertRaises(StateError):
            set_draft_status(self.db, draft, "sent", "send")
        self.assertEqual(self.db.query("SELECT status FROM drafts")[0]["status"], "approved")

    def test_invalid_contract_is_validation_error_not_crash(self):
        self.seed_company_contact_claim()
        for invalid in (None, [], "bad"):
            payload = self.draft_payload()
            payload["outreach_contract"] = invalid
            with self.assertRaises(ValidationError):
                create_draft(self.db, payload, "invalid-contract")

    def test_contact_address_change_and_identity_conflict_block_execution(self):
        self.seed_company_contact_claim()
        draft = create_draft(self.db, self.draft_payload(), "new")["draft_id"]
        approve_draft(self.db, draft, "human", "test", "approve")
        with self.db.transaction() as c:
            c.execute("UPDATE contacts SET public_email='other@example.test'")
        with self.assertRaises(StateError):
            set_draft_status(self.db, draft, "sent", "send")
        with self.db.transaction() as c:
            c.execute("UPDATE contacts SET public_email='alex@example.test',identity_status='contradicted'")
        with self.assertRaises(StateError):
            set_draft_status(self.db, draft, "sent", "send")

    def test_unaccepted_evidence_cannot_borrow_unrelated_accepted_claim(self):
        self.seed_company_contact_claim()
        with self.db.transaction() as c:
            c.execute("UPDATE claims SET status='proposed'")
        with self.assertRaises(StateError):
            create_draft(self.db, self.draft_payload(), "new")

    def test_expired_consent_and_catch_all_block_approval(self):
        self.seed_company_contact_claim()
        payload = self.draft_payload()
        payload["outreach_contract"]["consent"]["valid_until"] = "2020-01-01T00:00:00+00:00"
        draft = create_draft(self.db, payload, "new")["draft_id"]
        with self.assertRaises(StateError):
            approve_draft(self.db, draft, "human", "test", "approve")
        payload = self.draft_payload()
        payload["outreach_contract"]["address_status"] = "accept_all"
        draft = create_draft(self.db, payload, "next")["draft_id"]
        with self.assertRaises(StateError):
            approve_draft(self.db, draft, "human", "test", "approve")

    def test_one_channel_is_enough_and_failed_batch_rolls_back(self):
        self.seed_company_contact_claim()
        good = self.draft_payload()
        bad = {**good, "contact_id": "missing", "channel": "linkedin"}
        with self.assertRaises(ValidationError):
            ingest_draft_bundle(self.db, self.run, {"drafts": [good, bad]}, "batch")
        self.assertEqual(self.db.query("SELECT COUNT(*) AS n FROM drafts")[0]["n"], 0)
        result = ingest_draft_bundle(self.db, self.run, {"drafts": [good]}, "batch2")
        self.assertEqual(result["next_action"], "review_outreach")

    def test_draft_cannot_use_contact_from_other_run(self):
        self.seed_company_contact_claim()
        payload = self.draft_payload()
        payload["run_id"] = "another-run"
        with self.assertRaises(ValidationError):
            create_draft(self.db, payload, "new")

    def test_unknown_supplier_route_is_researchable_not_rejected(self):
        result = validate_company_fit_gates(
            {
                "name": "Test travel",
                "fit_status": "chat_candidate",
                "source_market_role": "outbound_seller",
                "china_product_signal": "adjacent",
                "ground_partner_route": "unclear",
                "competitive_position": "mixed",
            }
        )
        self.assertEqual(result["ground_partner_route"], "unclear")

    def test_arbitrary_email_prefix_is_not_personal_identity(self):
        self.assertIn("待确认", _route_kind("specialdesk@example.test", ""))

    def test_batch_export_contains_no_foreign_contacts_or_claims(self):
        self.seed_company_contact_claim()
        other = create_run(
            self.db, markets=["NZ"], segments=["travel_agency"], target=1, request_id="other", idempotency_key="other"
        )["run_id"]
        output = self.root / "other.json"
        export_json(self.db, output, other)
        tables = json.loads(output.read_text(encoding="utf-8"))["tables"]
        for table in ("contacts", "claims", "evidence", "claim_evidence", "drafts", "raw_imports"):
            self.assertEqual(tables[table], [], table)

    def test_single_sheet_links_and_formula_injection(self):
        from openpyxl import load_workbook

        self.seed_company_contact_claim()
        create_draft(self.db, self.draft_payload(), "new")
        with self.db.transaction() as c:
            c.execute("UPDATE companies SET name_original='=1+1'")
        output = self.root / "result.xlsx"
        export_workbook(self.db, output, self.run)
        wb = load_workbook(output)
        self.assertEqual(len(wb.sheetnames), 1)
        ws = wb.active
        self.assertEqual(ws["A2"].data_type, "s")
        self.assertEqual(ws.max_column, 7)
        self.assertIn("A6", ws["F2"].hyperlink.target)
        self.assertTrue(any("Would a regional outline help?" in str(c.value) for row in ws for c in row))
        wb.close()

    def test_migration_keeps_backup_and_legacy_data(self):
        self.seed_company_contact_claim()
        with self.db.transaction() as c:
            c.execute("UPDATE workspace_meta SET schema_version=3")
        result = self.db.migrate()
        from pathlib import Path

        self.assertTrue(Path(result["backup"]).is_file())
        self.assertEqual(self.db.query("SELECT COUNT(*) AS n FROM contacts")[0]["n"], 1)

    def test_install_detects_and_preserves_unknown_legacy_files(self):
        base = self.root / "skills"
        _skill_install(str(base), True)
        old = base / "inbound-ai-acquisition" / "legacy.py"
        old.write_text("user content", encoding="utf-8")
        self.assertFalse(_skill_status(str(base))["in_sync"])
        installed = _skill_install(str(base), True)
        from pathlib import Path

        self.assertEqual((Path(installed["backup"]) / "legacy.py").read_text(), "user content")
        self.assertTrue(installed["in_sync"])

    def test_missing_keys_never_call_network(self):
        with patch.dict(os.environ, {}, clear=True), patch("inbound_ai_acquisition.adapters.api.urlopen") as network:
            for provider in ("brave", "apollo", "hunter"):
                self.assertEqual(ApiAdapter(provider).discover({"query": "China tours"}).outcome, "unavailable")
            network.assert_not_called()

    def test_search_results_remain_clues_and_cost_unknown(self):
        response = BytesIO(
            json.dumps({"web": {"results": [{"url": "https://example.test", "description": "China"}]}}).encode()
        )
        with (
            patch.dict(os.environ, {"BRAVE_SEARCH_API_KEY": "test-secret"}),
            patch("inbound_ai_acquisition.adapters.api.urlopen", return_value=response),
        ):
            result = ApiAdapter("brave").discover({"query": "China tours"})
        self.assertTrue(result.observations[0]["requires_original_page"])
        self.assertIsNone(result.cost_amount)

    def test_api_rate_limit_is_explicit_and_secret_not_returned(self):
        with (
            patch.dict(os.environ, {"HUNTER_API_KEY": "test-secret"}),
            patch(
                "inbound_ai_acquisition.adapters.api.urlopen",
                side_effect=HTTPError("https://example.test?key=test-secret", 429, "secret", {}, None),
            ),
        ):
            result = ApiAdapter("hunter").inspect_contact({"email": "a@example.test"})
        self.assertEqual(result.outcome, "rate_limited")
        self.assertNotIn("test-secret", str(result))


if __name__ == "__main__":
    unittest.main()
