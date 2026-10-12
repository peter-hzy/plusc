"""Regressions for real-run defects; all records here are isolated fixtures."""

import copy
import unittest

import test_cli_workflow as fixtures
from test_integrated_workflow import fixture_bundle

from inbound_ai_acquisition.adapters.codex_bridge import CodexBridgeAdapter
from inbound_ai_acquisition.application.research import ingest_source_bundle, research_status
from inbound_ai_acquisition.errors import ValidationError
from inbound_ai_acquisition.exporters.workbook import export_workbook
from inbound_ai_acquisition.storage.database import utc_now
from inbound_ai_acquisition.storage.repositories import blocking_gap_count


class RealBatchRegressions(unittest.TestCase):
    setUp = fixtures.CliWorkflowTests.setUp
    tearDown = fixtures.CliWorkflowTests.tearDown

    def ingest(self, bundle):
        return ingest_source_bundle(self.db, self.run, bundle, "fixture-import")

    def test_email_supplement_reuses_contact_and_export_shows_it(self):
        bundle = fixture_bundle()
        self.ingest(bundle)
        cid = self.db.query("SELECT contact_id FROM contacts")[0]["contact_id"]
        bundle["companies"][0]["contact"]["public_email"] = "team@fixture.example.org"
        self.ingest(bundle)
        self.ingest(bundle)
        self.assertEqual(self.db.query("SELECT contact_id FROM contacts"), [{"contact_id": cid}])
        path = self.root / "supplement.xlsx"
        export_workbook(self.db, path, self.run)
        from openpyxl import load_workbook
        book = load_workbook(path)
        self.assertIn("team@fixture.example.org", book.active["E2"].value)
        book.close()

    def test_conflicting_email_and_unrelated_anonymous_page_stay_separate(self):
        bundle = fixture_bundle()
        bundle["companies"][0]["contact"]["public_email"] = "one@fixture.example.org"
        self.ingest(bundle)
        bundle["companies"][0]["contact"]["public_email"] = "two@fixture.example.org"
        self.ingest(bundle)
        self.assertEqual(research_status(self.db, self.run)["stats"]["contacts"], 2)
        bundle["companies"][0]["contact"].pop("public_email")
        bundle["companies"][0]["contact"]["sources"][0]["source_url"] += "/other"
        self.ingest(bundle)
        self.assertEqual(research_status(self.db, self.run)["stats"]["contacts"], 3)

    def test_known_email_match_wins_over_legacy_empty_handoff(self):
        bundle = fixture_bundle()
        self.ingest(bundle)
        contact = bundle["companies"][0]["contact"]
        contact["public_email"] = "team@fixture.example.org"
        original_source = copy.deepcopy(contact["sources"][0])
        contact["sources"][0]["source_url"] += "/other"
        self.ingest(bundle)
        contact["sources"].append(original_source)
        self.ingest(bundle)
        self.ingest(bundle)
        self.assertEqual(research_status(self.db, self.run)["stats"]["contacts"], 2)

    def test_anonymous_different_departments_on_same_page_do_not_merge(self):
        bundle = fixture_bundle()
        contact = bundle["companies"][0]["contact"]
        contact.update(route_grade="R2", route_purpose="business", why_contact="School partnerships")
        self.ingest(bundle)
        contact["why_contact"] = "Group partnerships"
        self.ingest(bundle)
        self.assertEqual(research_status(self.db, self.run)["stats"]["contacts"], 2)

    def test_export_prefers_usable_email_over_newer_empty_route(self):
        bundle = fixture_bundle()
        bundle["companies"][0]["contact"]["public_email"] = "team@fixture.example.org"
        self.ingest(bundle)
        bundle["companies"][0]["contact"].pop("public_email")
        bundle["companies"][0]["contact"]["sources"][0]["source_url"] += "/other"
        self.ingest(bundle)
        path = self.root / "old-email.xlsx"
        export_workbook(self.db, path, self.run)
        from openpyxl import load_workbook
        book = load_workbook(path)
        self.assertIn("team@fixture.example.org", book.active["E2"].value)
        book.close()

    def test_chinese_limits_get_conservative_wrapped_height(self):
        bundle = fixture_bundle()
        bundle["companies"][0]["risk_notes"] = "英文导游和大团接待能力待确认；" * 18
        self.ingest(bundle)
        path = self.root / "chinese.xlsx"
        export_workbook(self.db, path, self.run)
        from openpyxl import load_workbook
        book = load_workbook(path)
        self.assertIn("待确认", book.active["G2"].value)
        self.assertGreaterEqual(book.active.row_dimensions[2].height, 300)
        self.assertLessEqual(book.active.row_dimensions[2].height, 409)
        book.close()

    def test_product_gap_replays_and_cannot_be_silently_downgraded(self):
        bundle = fixture_bundle()
        company = bundle["companies"][0]
        company["research_gaps"] = [{"field_name": "group_capacity", "business_impact": "超过8人待确认",
                                     "blocking": True}]
        self.ingest(bundle)
        self.ingest(bundle)
        status = research_status(self.db, self.run)
        self.assertEqual(status["stats"]["research_gaps"], 1)
        with self.db.transaction() as conn:
            self.assertEqual(blocking_gap_count(conn, status["review"][0]["company"]["company_id"]), 1)
        company["research_gaps"][0]["blocking"] = False
        self.ingest(bundle)
        company.pop("research_gaps")
        self.ingest(bundle)
        self.assertEqual(self.db.query("SELECT blocking,status FROM research_gaps"),
                         [{"blocking": 1, "status": "open"}])
        self.assertEqual(self.db.query("SELECT COUNT(*) AS n FROM drafts")[0]["n"], 0)

    def test_bad_gap_boolean_rolls_back_all_changes(self):
        bundle = fixture_bundle()
        bundle["companies"][0]["research_gaps"] = [{"field_name": "guide", "business_impact": "待确认",
                                                     "blocking": "false"}]
        with self.assertRaises(ValidationError):
            self.ingest(bundle)
        self.assertEqual(self.db.query("SELECT COUNT(*) AS n FROM companies")[0]["n"], 0)

    def test_real_failed_attempts_are_visible_idempotent_and_not_evidence(self):
        bundle = fixture_bundle()
        bundle["source_attempts"] = [{"source_url": "https://fixture.example.org/restricted",
                                      "observed_at": utc_now(), "outcome": "rate_limited",
                                      "error_code": "HTTP_429", "retryable": True}]
        self.ingest(bundle)
        before = research_status(self.db, self.run)
        self.ingest(bundle)
        after = research_status(self.db, self.run)
        self.assertEqual(after["source_outcomes"]["rate_limited"], 1)
        self.assertEqual(before["stats"]["evidence"], after["stats"]["evidence"])
        row = self.db.query("SELECT * FROM source_runs WHERE outcome='rate_limited'")[0]
        self.assertEqual(row["result_count"], 0)
        self.assertEqual(row["started_at"], bundle["source_attempts"][0]["observed_at"])
        self.assertIn("restricted", row["query_summary"])

    def test_invalid_attempts_are_rejected(self):
        good = {"source_url": "https://fixture.example.org/blocked", "observed_at": utc_now(),
                "outcome": "restricted", "retryable": False}
        for change in ({"retryable": "false"}, {"observed_at": "2026-01-01"},
                       {"source_url": "file:///private"}, {"outcome": "succeeded"}, {"outcome": []}):
            bundle = fixture_bundle()
            bundle["source_attempts"] = [{**copy.deepcopy(good), **change}]
            with self.subTest(change=change), self.assertRaises(ValidationError):
                self.ingest(bundle)

    def test_bridge_rejects_failed_observations_and_string_boolean(self):
        for request in ({"outcome": "restricted", "observations": [{
            "source_url": "https://fixture.example.org/blocked", "observed_text": "access denied"}]},
                        {"outcome": "no_results", "retryable": "false"}):
            with self.assertRaises(ValidationError):
                CodexBridgeAdapter.inspect_company(request)
