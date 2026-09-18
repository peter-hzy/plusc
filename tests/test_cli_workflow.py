import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from inbound_ai_acquisition.adapters.manual import ManualAdapter
from inbound_ai_acquisition.adapters.registry import built_in_adapters
from inbound_ai_acquisition.application.discovery import discovery_coverage
from inbound_ai_acquisition.application.legacy_import import import_legacy_bundle
from inbound_ai_acquisition.application.research import (
    approve_run_claims,
    ingest_draft_bundle,
    ingest_source_bundle,
    research_status,
    verify_research_uat,
)
from inbound_ai_acquisition.application.services import (
    approve_draft,
    create_draft,
    create_run,
    import_companies,
    import_contacts,
    import_evidence,
    ingest_reply,
    record_source_result,
    set_draft_status,
)
from inbound_ai_acquisition.audit.public_repo import audit_public_repo
from inbound_ai_acquisition.cli import _context_validate, _init_workspace, _skill_install, _skill_status
from inbound_ai_acquisition.config import Workspace
from inbound_ai_acquisition.domain.policies import require_one_cta
from inbound_ai_acquisition.errors import StateError, ValidationError
from inbound_ai_acquisition.exporters.workbook import BUSINESS_SHEET, HEADERS, export_workbook
from inbound_ai_acquisition.storage.database import Database, utc_now


class CliWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.workspace = Workspace(self.root / "workspace")
        self.workspace.ensure_layout()
        self.db = Database(self.workspace.database)
        self.db.migrate()
        self.run = create_run(
            self.db,
            markets=["AU", "NZ"],
            segments=["travel_agency", "dmc"],
            target=20,
            request_id="req-run",
            idempotency_key="idem-run",
        )["run_id"]

    def tearDown(self):
        self.temp.cleanup()

    def write_json(self, name, payload):
        path = self.root / name
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        return path

    def seed_company_contact_claim(self):
        with self.db.transaction() as connection:
            connection.execute("UPDATE runs SET product_version='test-v1' WHERE run_id=?", (self.run,))
        company_file = self.write_json(
            "companies.json",
            [
                {
                    "company_id": "co_demo",
                    "name": "Example Travel",
                    "country": "AU",
                    "segment": "travel_agency",
                    "website": "https://example.test",
                    "fit_reason": "Offers tailored Asia trips",
                }
            ],
        )
        import_companies(self.db, self.run, company_file, "req-company")
        contact_file = self.write_json(
            "contacts.json",
            [
                {
                    "contact_id": "ct_demo",
                    "company_id": "co_demo",
                    "full_name": "Alex Example",
                    "title": "Product Manager",
                    "public_email": "alex@example.test",
                    "linkedin_url": "https://linkedin.com/in/alex-example",
                    "identity_status": "supported",
                    "identity_evidence_url": "https://example.test/team",
                    "identity_evidence_text": "The team page lists Alex as Product Manager.",
                    "observed_at": "2026-09-04T00:00:00+00:00",
                }
            ],
        )
        import_contacts(self.db, contact_file, "req-contact")
        evidence_file = self.write_json(
            "evidence.json",
            [
                {
                    "evidence_id": "ev_demo",
                    "entity_type": "company",
                    "entity_id": "co_demo",
                    "source_url": "https://example.test/china",
                    "observed_text": "The company publicly lists tailored trips to China.",
                    "source_level": "primary",
                }
            ],
        )
        import_evidence(self.db, evidence_file, "req-evidence")
        with self.db.transaction() as connection:
            connection.execute(
                "INSERT INTO claims VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    "cl_demo",
                    "company",
                    "co_demo",
                    "fit_reason",
                    "Offers tailored China trips",
                    "accepted",
                    "strong",
                    "",
                    "human",
                    utc_now(),
                    "官网证据",
                    None,
                    utc_now(),
                ),
            )
            connection.execute("INSERT INTO claim_evidence VALUES (?,?)", ("cl_demo", "ev_demo"))
            for field, value in {
                "source_market_role": "outbound_seller",
                "china_product_signal": "direct",
                "ground_partner_route": "external_partner_possible",
                "competitive_position": "potential_buyer",
            }.items():
                connection.execute(
                    "INSERT INTO claims VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        "cl_" + field,
                        "company",
                        "co_demo",
                        field,
                        value,
                        "accepted",
                        "strong",
                        "",
                        "test human",
                        utc_now(),
                        "isolated fixture",
                        None,
                        utc_now(),
                    ),
                )
                connection.execute("INSERT INTO claim_evidence VALUES (?,?)", ("cl_" + field, "ev_demo"))

    def test_run_create_is_idempotent(self):
        second = create_run(
            self.db,
            markets=["AU", "NZ"],
            segments=["travel_agency", "dmc"],
            target=20,
            request_id="req-run-retry",
            idempotency_key="idem-run",
        )
        self.assertEqual(second["status"], "already_applied")
        self.assertEqual(second["run_id"], self.run)

    def test_run_captures_product_and_voice_versions(self):
        result = create_run(
            self.db,
            markets=["AU"],
            segments=["travel_agency"],
            target=20,
            request_id="req-versioned-run",
            idempotency_key="idem-versioned-run",
            product_version="product-1.2",
            voice_version="voice-2.0",
        )
        row = self.db.query("SELECT product_version,voice_version FROM runs WHERE run_id=?", (result["run_id"],))[0]
        self.assertEqual(row["product_version"], "product-1.2")
        self.assertEqual(row["voice_version"], "voice-2.0")

    def test_unknown_import_field_rolls_back_entire_file(self):
        source = self.write_json(
            "bad.json",
            [
                {"name": "Good Before Bad", "website": "https://good.test"},
                {"name": "Bad", "website": "https://bad.test", "invented": "must fail"},
            ],
        )
        with self.assertRaises(ValidationError):
            import_companies(self.db, self.run, source, "req-bad")
        self.assertEqual(self.db.query("SELECT COUNT(*) AS n FROM companies")[0]["n"], 0)

    def test_verified_contact_without_identity_evidence_is_rejected(self):
        source = self.write_json(
            "company.json", [{"company_id": "co_one", "name": "One", "website": "https://one.test"}]
        )
        import_companies(self.db, self.run, source, "req-company")
        contacts = self.write_json(
            "bad-contact.json",
            [
                {
                    "company_id": "co_one",
                    "full_name": "Unproven Person",
                    "identity_status": "verified",
                }
            ],
        )
        with self.assertRaises(ValidationError):
            import_contacts(self.db, contacts, "req-bad-contact")
        self.assertEqual(self.db.query("SELECT COUNT(*) AS n FROM contacts")[0]["n"], 0)

    def test_approved_draft_can_be_marked_sent_but_never_sends(self):
        self.seed_company_contact_claim()
        result = create_draft(
            self.db,
            {
                "run_id": self.run,
                "contact_id": "ct_demo",
                "channel": "email",
                "subject": "China ground support",
                "body": "Hi Alex, I noticed your tailored China trips. Would a sample ground-service outline be useful?",
                "evidence_ids": ["ev_demo"],
                "product_version": "test-v1",
                "outreach_contract": {
                    "recipient_kind": "person",
                    "address_status": "published",
                    "role_reason": "Product manager",
                    "fact_checked_by": "test reviewer",
                    "product_approved_by": "test owner",
                    "consent": {
                        "basis": "express",
                        "market": "AU",
                        "source_reference": "test-only consent",
                        "reason": "isolated test",
                        "reviewed_by": "test reviewer",
                        "valid_until": "2099-01-01T00:00:00+00:00",
                    },
                },
            },
            "req-draft",
        )
        draft_id = result["draft_id"]
        approve_draft(self.db, draft_id, "human", "内容已核验", "req-approve")
        sent = set_draft_status(self.db, draft_id, "sent", "req-sent")
        self.assertEqual(sent["status"], "succeeded")
        row = self.db.query("SELECT status,sent_at FROM drafts WHERE draft_id=?", (draft_id,))[0]
        self.assertEqual(row["status"], "sent")
        self.assertIsNotNone(row["sent_at"])

    def test_unsubscribe_is_atomic_and_idempotent(self):
        self.seed_company_contact_claim()
        draft_id = create_draft(
            self.db,
            {
                "run_id": self.run,
                "contact_id": "ct_demo",
                "channel": "linkedin",
                "product_version": "test-v1",
                "body": "Hi Alex, I saw your China itineraries. Open to a quick ground-service outline?",
                "evidence_ids": ["ev_demo"],
            },
            "req-draft",
        )["draft_id"]
        payload = {
            "message_id": "msg-1",
            "contact_id": "ct_demo",
            "channel": "email",
            "received_at": "2026-09-04T00:00:00+00:00",
            "reply_category": "unsubscribe",
        }
        first = ingest_reply(self.db, payload, "reply-key", "req-reply")
        second = ingest_reply(self.db, payload, "reply-key", "req-reply-retry")
        self.assertGreaterEqual(first["changed_records"], 3)
        self.assertEqual(second["status"], "already_applied")
        self.assertEqual(self.db.query("SELECT COUNT(*) AS n FROM replies")[0]["n"], 1)
        self.assertEqual(self.db.query("SELECT COUNT(*) AS n FROM stop_rules WHERE status='active'")[0]["n"], 1)
        self.assertEqual(
            self.db.query("SELECT status FROM drafts WHERE draft_id=?", (draft_id,))[0]["status"], "stopped"
        )

    def test_workbook_round_trip_has_all_review_sheets(self):
        output = self.root / "batch.xlsx"
        result = export_workbook(self.db, output, self.run)
        self.assertEqual(result["sheets"], 1)
        self.assertTrue(output.exists())
        from openpyxl import load_workbook

        workbook = load_workbook(output, read_only=True)
        self.assertEqual(workbook.sheetnames[0], BUSINESS_SHEET)
        self.assertEqual(workbook.sheetnames, [BUSINESS_SHEET])
        workbook.close()

    def test_legacy_bundle_is_preserved_and_mapped(self):
        other_root = self.root / "legacy-workspace"
        Workspace(other_root).ensure_layout()
        other_db = Database(other_root / "acquisition.db")
        other_db.migrate()
        fixture = Path(__file__).resolve().parents[1] / "examples" / "fictional-workbook.json"
        first = import_legacy_bundle(other_db, fixture, "req-legacy")
        second = import_legacy_bundle(other_db, fixture, "req-legacy-retry")
        self.assertEqual(first["status"], "succeeded")
        self.assertEqual(second["status"], "already_applied")
        self.assertEqual(other_db.query("SELECT COUNT(*) AS n FROM raw_imports")[0]["n"], 1)
        self.assertEqual(other_db.query("SELECT COUNT(*) AS n FROM companies")[0]["n"], 5)
        self.assertEqual(other_db.query("SELECT COUNT(*) AS n FROM contacts")[0]["n"], 5)
        self.assertEqual(other_db.query("SELECT COUNT(*) AS n FROM evidence")[0]["n"], 10)

    def test_database_integrity_and_event_sequences(self):
        self.assertEqual(self.db.query("PRAGMA integrity_check")[0]["integrity_check"], "ok")
        events = self.db.query("SELECT sequence FROM events WHERE run_id=? ORDER BY sequence", (self.run,))
        self.assertEqual([row["sequence"] for row in events], list(range(1, len(events) + 1)))

    def test_source_restricted_is_not_no_results(self):
        payload = {
            "run_id": self.run,
            "adapter_id": "public_web",
            "query_summary": "site restricted during public research",
            "outcome": "restricted",
            "error_code": "http_403",
            "retryable": False,
            "observations": [],
        }
        result = record_source_result(self.db, payload, "source-key", "req-source")
        self.assertEqual(result["status"], "succeeded")
        row = self.db.query("SELECT outcome,error_code,result_count FROM source_runs")[0]
        self.assertEqual(row, {"outcome": "restricted", "error_code": "http_403", "result_count": 0})

    def test_conflicting_claim_cannot_be_accepted(self):
        self.seed_company_contact_claim()
        with self.db.transaction() as connection:
            connection.execute("UPDATE evidence SET conflicts=1 WHERE evidence_id='ev_demo'")
            connection.execute(
                "INSERT INTO claims VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    "cl_conflict",
                    "company",
                    "co_demo",
                    "fit_reason",
                    "Conflicted",
                    "proposed",
                    "strong",
                    "source conflict",
                    None,
                    None,
                    "",
                    None,
                    utc_now(),
                ),
            )
            connection.execute("INSERT INTO claim_evidence VALUES (?,?)", ("cl_conflict", "ev_demo"))
        from inbound_ai_acquisition.application.services import decide_claim

        with self.assertRaises(StateError):
            decide_claim(self.db, "cl_conflict", "accepted", "human", "try", "req-conflict")

    def test_public_repo_audit_ignores_fixture_domains_but_finds_database(self):
        project = self.root / "repo"
        project.mkdir()
        (project / "README.md").write_text("contact demo@example.test", encoding="utf-8")
        (project / ".env.example").write_text("SEARCH_API_KEY=\nOTHER_SECRET=\n", encoding="utf-8")
        self.assertEqual(audit_public_repo(project)["status"], "succeeded")
        (project / "acquisition.db").write_bytes(b"private")
        self.assertEqual(audit_public_repo(project)["status"], "failed")

    def test_release_lock_excludes_build_caches(self):
        from inbound_ai_acquisition.audit.release import _portable_hash, build_lock

        project = self.root / "release-repo"
        (project / "src").mkdir(parents=True)
        (project / "src" / "main.py").write_text("print('ok')\n", encoding="utf-8")
        for relative in (
            "build/copied.py",
            ".ruff_cache/cache",
            ".pytest_cache/cache",
            ".mypy_cache/cache",
            "node_modules/pkg/index.js",
            "src/demo.egg-info/PKG-INFO",
            ".coverage",
        ):
            path = project / relative
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("generated", encoding="utf-8")
        lock_path = project / "skills.lock.json"
        payload = build_lock(project, lock_path)
        self.assertEqual([row["path"] for row in payload["files"]], ["src/main.py"])
        lf = project / "lf.txt"
        crlf = project / "crlf.txt"
        lf.write_bytes(b"one\ntwo\n")
        crlf.write_bytes(b"one\r\ntwo\r\n")
        self.assertEqual(_portable_hash(lf), _portable_hash(crlf))

    def test_release_lock_uses_git_index_and_excludes_private_local_files(self):
        from inbound_ai_acquisition.audit.release import build_lock

        project = self.root / "git-release-repo"
        project.mkdir()
        subprocess.run(["git", "init"], cwd=project, check=True, capture_output=True)
        (project / ".gitignore").write_text("uat/\ninternal.md\n", encoding="utf-8")
        (project / "README.md").write_text("public\n", encoding="utf-8")
        (project / "internal.md").write_text("private notes\n", encoding="utf-8")
        private = project / "uat" / "buyers.json"
        private.parent.mkdir()
        private.write_text('{"real": true}\n', encoding="utf-8")
        subprocess.run(["git", "add", ".gitignore", "README.md"], cwd=project, check=True, capture_output=True)
        payload = build_lock(project, project / "skills.lock.json")
        self.assertEqual([row["path"] for row in payload["files"]], [".gitignore", "README.md"])

    def test_skill_install_and_status_use_manifest_diff(self):
        base = self.root / "skills"
        self.assertEqual(_skill_status(str(base))["status"], "not_installed")
        installed = _skill_install(str(base), replace=True)
        self.assertTrue(installed["in_sync"])
        target_skill = base / "inbound-ai-acquisition" / "SKILL.md"
        target_skill.write_text(target_skill.read_text(encoding="utf-8") + "\nchanged\n", encoding="utf-8")
        self.assertFalse(_skill_status(str(base))["in_sync"])

    def test_context_validation_rejects_unfilled_profiles(self):
        root = self.root / "context-workspace"
        _init_workspace(root)
        workspace = Workspace(root)
        db = Database(workspace.database)
        unfilled = _context_validate(workspace, db)
        self.assertEqual(unfilled["status"], "needs_review")
        self.assertIn("product-profile.md", unfilled["profile_gaps"])
        self.assertIn("voice-profile.md", unfilled["profile_gaps"])

        (root / "product-profile.md").write_text(
            "# 产品资料卡\n\n版本：1.0\n- 可接待目的地：北京、上海、西安\n"
            "- 可提供服务：线路设计、交通、导游\n- 语言：中文、英文\n",
            encoding="utf-8",
        )
        (root / "voice-profile.md").write_text(
            "# 话术资料卡\n\n版本：1.0\n- 希望像谁在说话：熟悉中国地接的合作伙伴\n",
            encoding="utf-8",
        )
        self.assertEqual(_context_validate(workspace, db)["status"], "succeeded")

    def test_cli_emits_utf8_when_environment_requests_ascii(self):
        root = self.root / "unicode-workspace"
        _init_workspace(root)
        environment = os.environ.copy()
        environment["PYTHONIOENCODING"] = "ascii"
        result = subprocess.run(
            [
                sys.executable,
                "-m",
                "inbound_ai_acquisition.cli",
                "context",
                "validate",
                "--workspace",
                str(root),
            ],
            capture_output=True,
            check=False,
            env=environment,
        )
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8", errors="replace"))
        self.assertIn("可接待目的地", result.stdout.decode("utf-8"))

    def test_real_adapters_are_not_manual_aliases(self):
        adapters = built_in_adapters()
        self.assertIsInstance(adapters["manual"], ManualAdapter)
        for name in ("public_web", "codex_bridge", "directory_csv"):
            self.assertNotIsInstance(adapters[name], ManualAdapter)
        self.assertEqual(adapters["public_web"].discover({"query": "AU travel"}).outcome, "restricted")
        self.assertEqual(adapters["public_web"].inspect_company({"url": "bad"}).outcome, "failed_terminal")
        self.assertEqual(
            adapters["directory_csv"].discover({"path": str(self.root / "missing.csv")}).outcome,
            "failed_terminal",
        )
        with self.assertRaises(ValidationError):
            adapters["codex_bridge"].discover({"outcome": "failed", "observations": []})

    def test_schema_v2_restores_legacy_candidate_pool_from_event(self):
        now = utc_now()
        with self.db.transaction() as connection:
            connection.execute(
                "INSERT INTO source_runs VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    "src_legacy",
                    self.run,
                    "codex_bridge",
                    "Legacy discovery query",
                    "legacy-fingerprint",
                    now,
                    now,
                    1,
                    "succeeded",
                    2,
                    0.0,
                    0,
                    "",
                    0,
                ),
            )
            self.db.append_event(
                connection,
                event_type="research.sources_ingested",
                request_id="legacy-event",
                run_id=self.run,
                payload={
                    "discovered_candidates": [
                        {"name": "Legacy One", "source_url": "https://legacy-one.example.com"},
                        {"name": "Legacy Two", "source_url": "https://legacy-two.example.com"},
                    ]
                },
            )
            connection.execute("DROP TABLE candidate_sources")
            connection.execute("DROP TABLE candidates")
            connection.execute("DROP TABLE source_plan_entries")
            connection.execute("UPDATE workspace_meta SET schema_version=1 WHERE singleton=1")
        self.db.migrate()
        self.assertEqual(self.db.schema_version(), 5)
        self.assertEqual(len(self.db.query("SELECT * FROM candidates WHERE run_id=?", (self.run,))), 2)
        self.assertEqual(len(self.db.query("SELECT * FROM candidate_sources WHERE run_id=?", (self.run,))), 2)
        plan = self.db.query("SELECT * FROM source_plan_entries WHERE run_id=?", (self.run,))
        self.assertEqual(plan[0]["source_category"], "legacy_codex_discovery")

    def test_discovery_can_be_extended_and_deduplicates_by_domain(self):
        def bundle(query, category, rows):
            return {
                "executor": "codex_bridge",
                "discovered_count": len(rows),
                "discovery_query": query,
                "discovered_candidates": [
                    {
                        "name": name,
                        "source_url": url,
                        "website": url,
                        "observed_text": text,
                        "source_category": category,
                        "query_summary": query,
                    }
                    for name, url, text in rows
                ],
                "companies": [],
            }

        ingest_source_bundle(
            self.db,
            self.run,
            bundle(
                "first wave",
                "industry_directory",
                [
                    ("Alpha Travel", "https://alpha.example.org/china", "Alpha appeared in the first source."),
                    ("Beta Travel", "https://beta.example.org/china", "Beta appeared in the first source."),
                ],
            ),
            "wave-one",
        )
        result = ingest_source_bundle(
            self.db,
            self.run,
            bundle(
                "second wave",
                "public_web_search",
                [
                    ("Alpha Holidays", "https://alpha.example.org/about", "Same company under a brand variation."),
                    ("Gamma Travel", "https://gamma.example.org/china", "Gamma appeared in the second source."),
                ],
            ),
            "wave-two",
        )
        self.assertEqual(result["coverage"]["candidate_count"], 3)
        alpha = self.db.query(
            "SELECT candidate_id FROM candidates WHERE run_id=? AND root_domain=?",
            (self.run, "alpha.example.org"),
        )
        self.assertEqual(len(alpha), 1)
        sources = self.db.query(
            "SELECT source_category FROM candidate_sources WHERE candidate_id=?", (alpha[0]["candidate_id"],)
        )
        self.assertEqual({row["source_category"] for row in sources}, {"industry_directory", "public_web_search"})

    def test_schema_v3_merges_legacy_candidate_domain_duplicates(self):
        now = utc_now()
        with self.db.transaction() as connection:
            connection.execute("DROP INDEX IF EXISTS uq_candidate_run_domain")
            for candidate_id, name in (("cand_old_a", "Legacy Alpha"), ("cand_old_b", "Alpha Holidays")):
                connection.execute(
                    "INSERT INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        candidate_id,
                        self.run,
                        name,
                        name.casefold(),
                        "alpha.example.org",
                        f"https://alpha.example.org/{candidate_id}",
                        "AU",
                        "travel_agency",
                        "candidate_pool",
                        "legacy duplicate",
                        None,
                        now,
                        now,
                    ),
                )
            connection.execute("UPDATE workspace_meta SET schema_version=2 WHERE singleton=1")
        self.db.migrate()
        rows = self.db.query(
            "SELECT candidate_id FROM candidates WHERE run_id=? AND root_domain='alpha.example.org'", (self.run,)
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(self.db.schema_version(), 5)

    def test_run_discovery_policy_is_configurable_and_market_neutral(self):
        configured = create_run(
            self.db,
            markets=["NZ"],
            segments=["mice"],
            target=5,
            request_id="req-configured",
            idempotency_key="idem-configured",
            candidate_pool_multiplier=2,
            min_source_categories=2,
            production_min_target=5,
        )["run_id"]
        coverage = discovery_coverage(self.db, configured)
        self.assertEqual(coverage["recommended_pool"], 10)
        self.assertEqual(coverage["policy"]["min_source_categories"], 2)
        verification = verify_research_uat(self.db, configured)
        self.assertFalse(any("AU" in item for item in verification["failures"]))

    def test_same_domain_can_be_researched_in_different_runs(self):
        source = self.write_json(
            "repeat-company.json", [{"name": "Repeat Travel", "website": "https://repeat.example.org"}]
        )
        import_companies(self.db, self.run, source, "repeat-one")
        second_run = create_run(
            self.db,
            markets=["NZ"],
            segments=["travel_agency"],
            target=5,
            request_id="repeat-run",
            idempotency_key="repeat-run",
        )["run_id"]
        import_companies(self.db, second_run, source, "repeat-two")
        self.assertEqual(
            self.db.query("SELECT COUNT(*) AS n FROM companies WHERE root_domain='repeat.example.org'")[0]["n"],
            2,
        )

    def test_research_bridge_pauses_for_human_then_creates_two_native_drafts(self):
        with self.db.transaction() as connection:
            connection.execute("UPDATE runs SET product_version='test-v1' WHERE run_id=?", (self.run,))
        bundle = {
            "executor": "codex_bridge",
            "discovered_count": 2,
            "discovery_query": "Australian China tour operators",
            "source_plan": [
                {
                    "source_category": "industry_directory",
                    "source_role": "discover",
                    "target_signal": "Australian tour operator",
                    "query_or_entry_url": "https://example.com/directory",
                    "coverage": "AU tour operators",
                    "access_status": "opened",
                },
                {
                    "source_category": "public_web_search",
                    "source_role": "discover",
                    "target_signal": "China product",
                    "query_or_entry_url": "Australian China tour operators",
                    "coverage": "AU China products",
                    "access_status": "opened",
                },
            ],
            "discovered_candidates": [
                {
                    "name": "Example Operator",
                    "source_url": "https://example.com/directory",
                    "observed_text": "Public search result identified Example Operator as an Australian China tour operator.",
                    "source_category": "industry_directory",
                    "query_summary": "Australian tour operator directory",
                    "country_hint": "AU",
                    "segment_hint": "tour_operator",
                    "sources": [
                        {
                            "source_url": "https://example.com/china",
                            "observed_text": "The operator also appears through a China product search.",
                            "source_category": "public_web_search",
                            "query_summary": "Australian China tour operators",
                        }
                    ],
                },
                {
                    "name": "Alternative Operator",
                    "source_url": "https://alternative.example.org/china",
                    "observed_text": "Public result identified another Australian operator for screening.",
                    "source_category": "public_web_search",
                    "query_summary": "Australian China tour operators",
                    "country_hint": "AU",
                    "segment_hint": "travel_agency",
                    "decision_reason": "Needs official-site verification before selection.",
                },
            ],
            "companies": [
                {
                    "name": "Example Operator",
                    "country": "AU",
                    "city": "Sydney",
                    "segment": "tour_operator",
                    "website": "https://example.com",
                    "fit_reason": "Publishes tailored China itineraries.",
                    "fit_status": "priority_chat",
                    "source_market_role": "outbound_seller",
                    "china_product_signal": "direct",
                    "ground_partner_route": "external_partner_possible",
                    "competitive_position": "potential_buyer",
                    "sources": [
                        {
                            "operation": "inspect_company",
                            "query_summary": "Open official China itinerary",
                            "source_url": "https://example.com/china",
                            "observed_text": "Official website describes tailored China itineraries.",
                            "source_level": "primary",
                            "purpose": "company_fit",
                        }
                    ],
                    "contact": {
                        "full_name": "Partnerships Team",
                        "title": "Partnerships",
                        "why_contact": "Official business handoff route for supplier enquiries.",
                        "contact_role": "handoff",
                        "public_email": "partners@example.com",
                        "identity_status": "supported",
                        "sources": [
                            {
                                "query_summary": "Open official contact page",
                                "source_url": "https://example.com/contact",
                                "observed_text": "Official contact page lists partners@example.com for partnership enquiries.",
                                "source_level": "primary",
                                "purpose": "contact_route",
                            }
                        ],
                    },
                }
            ],
        }
        result = ingest_source_bundle(self.db, self.run, bundle, "req-research")
        # 待批准公司正式域名尚为空，也必须沿候选实体链接去重。
        ingest_source_bundle(self.db, self.run, bundle, "req-research-repeat")
        self.assertEqual(len(self.db.query("SELECT * FROM companies")), 1)
        self.assertEqual(result["next_action"], "awaiting_claim_review")
        self.assertEqual(result["coverage"]["candidate_count"], 2)
        self.assertEqual(result["coverage"]["selected_count"], 1)
        self.assertEqual(result["coverage"]["source_category_count"], 2)
        performance = {row["source_category"]: row for row in result["coverage"]["source_performance"]}
        self.assertEqual(performance["industry_directory"]["unique_candidates"], 0)
        self.assertEqual(performance["industry_directory"]["overlap_candidates"], 1)
        self.assertFalse(result["coverage"]["production_ready"])
        before = research_status(self.db, self.run)
        self.assertGreater(before["stats"]["evidence"], 0)
        self.assertGreater(before["stats"]["proposed_claims"], 0)
        self.assertEqual(before["review"][0]["decision_state"], "待人工批准")
        self.assertEqual(before["review"][0]["company_preview"]["country"], "AU")
        self.assertEqual(before["review"][0]["contact_previews"][0]["full_name"], "Partnerships Team")
        self.assertEqual(len(before["candidates"]), 2)
        self.assertEqual({row["discovery_status"] for row in before["candidates"]}, {"candidate_pool", "selected"})
        approved = approve_run_claims(self.db, self.run, "test human", "Reviewed sources", "req-human")
        self.assertEqual(approved["next_action"], "awaiting_drafts")
        contact = self.db.query("SELECT contact_id FROM contacts")[0]
        evidence = self.db.query(
            "SELECT e.evidence_id FROM evidence e JOIN claim_evidence ce ON ce.evidence_id=e.evidence_id JOIN claims cl ON cl.claim_id=ce.claim_id WHERE e.entity_type='company' AND cl.status='accepted'"
        )[0]
        draft_bundle = {
            "drafts": [
                {
                    "product_version": "test-v1",
                    "contact_id": contact["contact_id"],
                    "channel": "email",
                    "subject": "China ground support",
                    "body": "Hi, I noticed your tailored China itineraries. We support custom ground arrangements in China. Would a short capability outline be useful?",
                    "evidence_ids": [evidence["evidence_id"]],
                },
                {
                    "product_version": "test-v1",
                    "contact_id": contact["contact_id"],
                    "channel": "linkedin",
                    "body": "Hi, I saw your tailored China trips and wanted to connect. Would it help if I shared a concise outline of our local ground support?",
                    "evidence_ids": [evidence["evidence_id"]],
                },
            ]
        }
        finished = ingest_draft_bundle(self.db, self.run, draft_bundle, "req-drafts")
        self.assertEqual(finished["next_action"], "review_outreach")
        after = research_status(self.db, self.run)
        self.assertEqual(after["stats"]["review_ready_email"], 1)
        self.assertEqual(after["stats"]["review_ready_linkedin"], 1)

        output = self.root / "uat.xlsx"
        export_workbook(self.db, output, self.run)
        from openpyxl import load_workbook

        opened = load_workbook(output, read_only=False)
        lead_review = opened[BUSINESS_SHEET]
        self.assertEqual([cell.value for cell in lead_review[1]], HEADERS)
        self.assertEqual(lead_review.auto_filter.ref, "A1:G2")
        self.assertIn("企业转交邮箱", lead_review["E2"].value)
        self.assertGreater(lead_review.max_row, 2)
        self.assertEqual(opened.sheetnames, [BUSINESS_SHEET])
        opened.close()
        verification = verify_research_uat(self.db, self.run, output)
        self.assertFalse(verification["passed"])
        self.assertNotIn("Excel 首页总览缺少业务审阅必需列", verification["failures"])
        self.assertIn("公司数不足：1/20", verification["failures"])
        self.assertFalse(verification["production_ready"])
        self.assertTrue(verification["coverage_warnings"])

    def test_draft_requires_exactly_one_cta(self):
        require_one_cta("Could I send a short outline?")
        with self.assertRaises(ValidationError):
            require_one_cta("This message has no ask.")
        with self.assertRaises(ValidationError):
            require_one_cta("Would you like the outline? Let me know.")

    def test_selected_company_requires_buy_side_fit_gates(self):
        from inbound_ai_acquisition.application.research import validate_company_fit_gates

        with self.assertRaisesRegex(ValidationError, "客源市场角色"):
            validate_company_fit_gates(
                {
                    "name": "Wrong Way DMC",
                    "fit_status": "priority_chat",
                    "china_product_signal": "direct",
                    "ground_partner_route": "unclear",
                    "competitive_position": "unknown",
                }
            )

    def test_competitor_cannot_be_priority_candidate(self):
        from inbound_ai_acquisition.application.research import validate_company_fit_gates

        with self.assertRaisesRegex(ValidationError, "竞争者"):
            validate_company_fit_gates(
                {
                    "name": "Integrated China Operator",
                    "fit_status": "priority_chat",
                    "source_market_role": "outbound_seller",
                    "china_product_signal": "direct",
                    "ground_partner_route": "in_house",
                    "competitive_position": "competitor",
                }
            )


if __name__ == "__main__":
    unittest.main()
