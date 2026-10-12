"""Synthetic service-fit fixtures only; never real leads or approvals."""

import copy
import unittest

import test_cli_workflow as fixtures
import test_release_behavior as release_fixtures
from test_integrated_workflow import fixture_bundle

from inbound_ai_acquisition.application.research import (
    ingest_source_bundle,
    research_status,
    verify_company_list,
)
from inbound_ai_acquisition.application.services import approve_draft, create_draft, create_run
from inbound_ai_acquisition.cli import build_parser
from inbound_ai_acquisition.domain.company_fit import SERVICE_FIELDS, validate_service_fit
from inbound_ai_acquisition.errors import StateError, ValidationError
from inbound_ai_acquisition.exporters.workbook import export_workbook
from inbound_ai_acquisition.storage.database import utc_now


def service_bundle():
    bundle = fixture_bundle()
    c = bundle["companies"][0]
    for field in ("source_market_role", "china_product_signal", "ground_partner_route", "competitive_position"):
        c.pop(field)
    c.update(
        segment="engineering", qualification_policy="service_fit_v1", fit_basis="supported",
        business_signal="Synthetic fixture: company announces a visiting team.",
        service_fit="Synthetic fixture: service location and visit overlap; product pending.",
        cooperation_hypothesis="Optional day activity, need not confirmed",
        counter_evidence="Attendee origin and free time remain unknown",
        next_action="Confirm arrangements with event organizer", fit_status="priority_chat",
    )
    url = c["sources"][0]["source_url"]
    c["field_sources"] = {f: [url] for f in (
        "website", "country", "segment", "fit_reason", "fit_status", "qualification_policy", *SERVICE_FIELDS,
    )}
    return bundle


class UnifiedPoolTests(unittest.TestCase):
    setUp = fixtures.CliWorkflowTests.setUp
    tearDown = fixtures.CliWorkflowTests.tearDown
    write_json = fixtures.CliWorkflowTests.write_json
    seed_company_contact_claim = fixtures.CliWorkflowTests.seed_company_contact_claim
    draft_payload = release_fixtures.ReleaseBehaviorTests.draft_payload

    def run_bundle(self, bundle=None, segments=None):
        run = create_run(self.db, markets=["AU"], segments=segments or ["*"], target=1,
                         request_id="unified", idempotency_key="unified")["run_id"]
        ingest_source_bundle(self.db, run, bundle or service_bundle(), "import")
        return run

    def test_cli_defaults_to_open_industry(self):
        args = build_parser().parse_args([
            "run", "create", "--workspace", str(self.root), "--market", "AU", "--research-only",
        ])
        self.assertEqual(args.segments, "*")

    def test_non_travel_company_roundtrip_and_replay(self):
        run = self.run_bundle()
        before = research_status(self.db, run)
        ingest_source_bundle(self.db, run, service_bundle(), "replay")
        after = research_status(self.db, run)
        for key in ("companies", "contacts", "evidence", "claims"):
            self.assertEqual(before["stats"][key], after["stats"][key])
        path = self.root / "result.xlsx"
        export_workbook(self.db, path, run)
        result = verify_company_list(self.db, run, path)
        self.assertTrue(result["passed"], result)
        from openpyxl import load_workbook
        book = load_workbook(path)
        self.assertEqual(len(book.sheetnames), 1)
        self.assertEqual(book.active.max_column, 7)
        self.assertIn("visiting team", book.active["C2"].value)
        book.close()

    def test_explicit_industry_restriction_still_applies(self):
        run = self.run_bundle(segments=["travel_agency"])
        result = verify_company_list(self.db, run)
        self.assertFalse(result["passed"])
        self.assertTrue(any("客户类型" in f for f in result["failures"]))

    def test_missing_rationale_and_invalid_types_fail(self):
        c = service_bundle()["companies"][0]
        for field in SERVICE_FIELDS:
            for value in ("", None, False, []):
                with self.subTest(field=field, value=value), self.assertRaises(ValidationError):
                    validate_service_fit({**c, field: value})
        with self.assertRaises(ValidationError):
            validate_service_fit({**c, "fit_basis": "hypothesis"})

    def test_missing_or_unrelated_binding_rolls_back(self):
        for change in (None, "foreign"):
            bundle = service_bundle()
            if change is None:
                bundle["companies"][0].pop("field_sources")
            else:
                bundle["companies"][0]["field_sources"]["business_signal"] = ["https://other.example.org"]
            with self.assertRaises(ValidationError):
                ingest_source_bundle(self.db, self.run, bundle, str(change))
            self.assertEqual(self.db.query("SELECT COUNT(*) AS n FROM companies")[0]["n"], 0)

    def seed_service_claims(self):
        self.seed_company_contact_claim()
        c = service_bundle()["companies"][0]
        with self.db.transaction() as conn:
            conn.execute("UPDATE companies SET segment='engineering' WHERE company_id='co_demo'")
            conn.execute("UPDATE claims SET proposed_value='not_source_market' WHERE field_name='source_market_role'")
            for field in ("qualification_policy", *SERVICE_FIELDS):
                conn.execute("INSERT INTO claims VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)", (
                    "unified_" + field, "company", "co_demo", field, c[field], "accepted", "strong",
                    "", "test reviewer", utc_now(), "synthetic only", None, utc_now(),
                ))
                conn.execute("INSERT INTO claim_evidence VALUES (?,?)", ("unified_" + field, "ev_demo"))

    def test_non_travel_outreach_uses_same_safety_gates(self):
        self.seed_service_claims()
        draft = create_draft(self.db, self.draft_payload(), "draft")["draft_id"]
        approve_draft(self.db, draft, "test reviewer", "synthetic only", "approve")
        self.assertEqual(self.db.query("SELECT status FROM drafts WHERE draft_id=?", (draft,))[0]["status"], "approved")

    def test_no_product_or_no_accepted_fit_does_not_bypass(self):
        self.seed_service_claims()
        payload = copy.deepcopy(self.draft_payload())
        payload["product_version"] = "unapproved-version"
        with self.assertRaises(StateError):
            create_draft(self.db, payload, "wrong-product")
        with self.db.transaction() as conn:
            conn.execute("UPDATE claims SET status='proposed' WHERE field_name='service_fit'")
        draft = create_draft(self.db, self.draft_payload(), "draft")["draft_id"]
        with self.assertRaises(StateError):
            approve_draft(self.db, draft, "test reviewer", "synthetic only", "approve")

    def test_conflicted_fit_source_blocks_approval(self):
        self.seed_service_claims()
        draft = create_draft(self.db, self.draft_payload(), "draft")["draft_id"]
        with self.db.transaction() as conn:
            conn.execute("UPDATE evidence SET conflicts=1 WHERE evidence_id='ev_demo'")
        with self.assertRaises(StateError):
            approve_draft(self.db, draft, "test reviewer", "synthetic only", "approve")

    def test_pending_contract_cannot_fall_back_to_old_buyer_approval(self):
        self.seed_service_claims()
        with self.db.transaction() as conn:
            conn.execute("UPDATE claims SET status='proposed' WHERE field_name='qualification_policy'")
            conn.execute("UPDATE claims SET proposed_value='outbound_seller' WHERE field_name='source_market_role'")
        draft = create_draft(self.db, self.draft_payload(), "draft")["draft_id"]
        with self.assertRaises(StateError):
            approve_draft(self.db, draft, "test reviewer", "synthetic only", "approve")
